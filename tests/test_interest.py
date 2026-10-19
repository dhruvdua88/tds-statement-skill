#!/usr/bin/env python3
"""The interest and late-fee arithmetic, on hand-checkable figures."""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import codes  # noqa: E402
import excel  # noqa: E402
import interest  # noqa: E402
import lint  # noqa: E402
from model import Statement  # noqa: E402

FIX = ROOT / "tests" / "fixtures"


class TestMonths(unittest.TestCase):
    def test_part_month_is_a_full_month(self):
        self.assertEqual(interest.months_or_part(date(2026, 7, 31), date(2026, 8, 1)), 2)
        self.assertEqual(interest.months_or_part(date(2026, 7, 1), date(2026, 7, 2)), 1)
        self.assertEqual(interest.months_or_part(date(2026, 12, 15), date(2027, 2, 1)), 3)
        self.assertEqual(interest.months_or_part(date(2026, 7, 5), date(2026, 7, 5)), 0)

    def test_deposit_due_dates(self):
        self.assertEqual(interest.deposit_due(date(2026, 7, 31)), date(2026, 8, 7))
        self.assertEqual(interest.deposit_due(date(2026, 12, 3)), date(2027, 1, 7))
        self.assertEqual(interest.deposit_due(date(2027, 3, 31)), date(2027, 4, 30))

    def test_statement_due_dates_come_from_one_table(self):
        self.assertEqual(set(codes.STATEMENT_DUE), {"Q1", "Q2", "Q3", "Q4"})
        self.assertEqual(interest.statement_due("2026-27", "Q2"), date(2026, 10, 31))
        self.assertEqual(interest.statement_due("2026-27", "Q3"), date(2027, 1, 31))
        self.assertEqual(interest.statement_due("2026-27", "Q4"), date(2027, 5, 31))


class TestStatement(unittest.TestCase):
    def st(self):
        return Statement.load(FIX / "example_140.json")

    def test_on_time_statement_owes_nothing(self):
        w = interest.compute(self.st(), filed_on="30-10-2026")
        self.assertEqual((w["interest"], w["fee"], w["days_late"]), (Decimal("0.00"), Decimal("0.00"), 0))

    def test_late_deduction_and_late_deposit(self):
        st = self.st()
        # row 1: credited 31-07, deducted 02-09 -> Jul, Aug, Sep = 3 months at 1% on 15000
        st.deductees[0].deduction_date = "02-09-2026"
        # challan 1 (rows 1 and 2) paid 08-09 instead of 07-08: row 2 deducted 31-07,
        # due 07-08, so 31-07 to 08-09 = Jul, Aug, Sep = 3 months at 1.5% on 30000;
        # row 1 deducted 02-09 is due 07-10, so it is on time.
        st.challans[0].challan_date = "08-09-2026"
        w = interest.compute(st, filed_on="30-10-2026")
        r1, r2 = w["rows"][0], w["rows"][1]
        self.assertEqual((r1["late_deduction_months"], r1["late_deduction_interest"]), (3, Decimal("450.00")))
        self.assertEqual(r1["late_deposit_interest"], Decimal("0.00"))
        self.assertEqual((r2["late_deposit_months"], r2["late_deposit_interest"]), (3, Decimal("1350.00")))
        self.assertEqual(w["interest"], Decimal("1800.00"))
        self.assertEqual(w["interest_short"], Decimal("1785.00"), "challan 3 already carries 15")
        self.assertIn("INTEREST_SHORT", {i.code for i in lint.model_issues(st)})
        self.assertEqual(st.challans[0].interest, Decimal("0.00"), "never added to a challan")

    def test_late_fee_per_day_and_the_cap(self):
        w = interest.compute(self.st(), filed_on="10-11-2026")
        self.assertEqual((w["days_late"], w["fee"]), (10, Decimal("2000.00")))
        w = interest.compute(self.st(), filed_on="31-10-2028")
        self.assertEqual(w["fee_cap"], Decimal("88500.00"))
        self.assertEqual(w["fee"], Decimal("88500.00"))
        self.assertGreater(w["fee_uncapped"], w["fee"])

    def test_the_workbook_carries_the_sheet(self):
        from openpyxl import load_workbook
        out = Path(tempfile.mkdtemp()) / "w.xlsx"
        excel.build(self.st(), out, filed_on="10-11-2026")
        ws = load_workbook(out)["Interest & Late Fee"]
        text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        self.assertIn("s.427 [s.234E of the 1961 Act]", text)
        self.assertIn("s.398(3)(a)(ii) [s.201(1A)(ii) of the 1961 Act]", text)
        self.assertIn("[verify", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
