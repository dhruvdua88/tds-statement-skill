#!/usr/bin/env python3
"""Form 138 layout tests: what the writer produces, and the workbook it follows.

    python3 -m unittest discover -s tests -p 'test_layout138.py' -v

Three levels of proof, and one honest limit. The first class checks the layout
module's own declarations. The fixture class checks the file the writer produces
from the synthetic salary statement: record widths, the head values, the global
deductee serial, file hygiene, the merged tax fields and the deposit date. The
workbook class re-reads the department's file-format workbook and holds every
field number in the layout maps against the official field name at that number,
so a field number that drifts from the authority fails here rather than at
upload; it skips when the workbook is not on this machine. None of this can
prove the utility will accept the file - only a run of the official FVU can do
that.
"""

from __future__ import annotations

import json
import re
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

import build_txt  # noqa: E402
import layout138  # noqa: E402
import lint  # noqa: E402
from model import Statement  # noqa: E402

FIX = ROOT / "tests" / "fixtures"

try:
    import openpyxl
except ImportError:  # pragma: no cover - the workbook class skips without it
    openpyxl = None


# ---------------------------------------------------------------------------
# The workbook cross-check
# ---------------------------------------------------------------------------

WORKBOOK_CANDIDATES = [
    Path.home() / ".tds-utils" / "format",
    Path.home() / "TDS-Skill-Research" / "official",
    Path.home() / "TDS-Skill-Research" / "downloads",
    Path.home() / "Downloads",
]

# Every field the layout maps, against the official field name at that number.
# The names below are transcribed from the Form 138 Q1-Q3 workbook once; the
# numbers are read from the layout maps and looked up in the workbook, so a
# number that has moved is caught. Keep the two in step when either changes.
EXPECTED_NAMES = {
    ("FH", "file_type"): "File Type",
    ("FH", "upload_type"): "Upload Type",
    ("FH", "file_creation_date"): "File Creation Date",
    ("FH", "file_sequence_no"): "File Sequence No.",
    ("FH", "uploader_type"): "Uploader Type",
    ("FH", "tan"): "TAN of Deductor",
    ("FH", "total_batches"): "Total No. of Batches",
    ("FH", "utility_name"): "Name of Return Preparation Utility",

    ("BH", "batch_number"): "Batch Number",
    ("BH", "num_challans"): "Count of Challan/transfer voucher Records",
    ("BH", "form_number"): "Form Number",
    ("BH", "token_previous"): "Token no. of previous regular statement (Form no. 138).",
    ("BH", "tan"): "TAN of Employer",
    ("BH", "pan_deductor"): "PAN of Deductor",
    ("BH", "assessment_year"): "Assessment Yr",
    ("BH", "tax_year"): "Tax Year",
    ("BH", "period"): "Period",
    ("BH", "name_deductor"): "Name of Deductor",
    ("BH", "addr_country"): "Deductor Address 6 (Country/Region)",
    ("BH", "addr1"): "Deductor Address1 (Flat/Door/Block number)",
    ("BH", "addr2"): "Deductor Address2 (Post Office)",
    ("BH", "addr3"): "Deductor Address3 (Road/Street/Block/Sector)",
    ("BH", "addr4"): "Deductor Address4 (Area/locality)",
    ("BH", "addr5"): "Deductor Address5 (District)",
    ("BH", "state_code"): "Deductor's Address - State",
    ("BH", "pin"): "Deductor's Address - Pincode",
    ("BH", "deductor_email"): "Deductor Email ID",
    ("BH", "deductor_country_code"): "Deductor Contact Country Code",
    ("BH", "deductor_phone"): "Deductor Contact Number",
    ("BH", "deductor_category"): "Deductor Type",
    ("BH", "rp_name"): "Name of Person responsible for Deduction",
    ("BH", "rp_designation"): "Designation of the Person responsible for Deduction",
    ("BH", "rp_addr1"): "Responsible Person's Address1 (Flat/Door/Block number)",
    ("BH", "rp_addr2"): "Responsible Person's Address2 (Post Office)",
    ("BH", "rp_addr3"): "Responsible Person's Address3 (Road/Street/Block/Sector)",
    ("BH", "rp_addr4"): "Responsible Person's Address4 (Area/locality)",
    ("BH", "rp_addr5"): "Responsible Person's Address5 (District)",
    ("BH", "rp_state_code"): "Responsible Person's State",
    ("BH", "rp_pin"): "Responsible Person's PIN",
    ("BH", "rp_email"): "Responsible Person's Email ID -1",
    ("BH", "rp_country"): "Responsible Persons country region",
    ("BH", "rp_country_code"): "Responsible Person Contact Country/Region",
    ("BH", "rp_phone"): "Responsible Person Contact Number",
    ("BH", "batch_total"): "Batch Total of - Total of Deposit Amount as per Challan",
    ("BH", "salary_detail_count"): "Count of Salary Details Records (Not applicable)",
    ("BH", "regular_earlier"): "Whether regular statement for Form 138 filed for earlier period",
    ("BH", "state_name"): "State Name",
    ("BH", "ministry_code"): "Ministry Name",
    ("BH", "ministry_name"): "Ministry Name (Others)",
    ("BH", "rp_pan"): "PAN of Responsible Person",
    ("BH", "ain"): "Account Office Identification Number (AIN) of PAO/ TO/ CDDO",
    ("BH", "gstin"): "Goods and Service Tax Number (GSTN)",

    ("CD", "batch_number"): "Batch Number",
    ("CD", "challan_no"): "Challan-Detail Record Number (Sl no. of the form)",
    ("CD", "num_deductees"): "Count of Employee/Deductee / Party Records",
    ("CD", "nil_challan"): "NIL Challan Indicator",
    ("CD", "tax"): "Total tax Deducted",
    ("CD", "interest"): "Total Interest",
    ("CD", "fee"): "Total fee",
    ("CD", "penalty"): "Total Penalty",
    ("CD", "total"): ("Total of Deposit Amount as per Challan/Total Amount adjusted through "
                      "Book Adjustment (Total of values in B+C+D+E)"),
    ("CD", "challan_mode"): "Mode of payment of tax (by way of challan/by way of book adjustment)",
    ("CD", "bsr_code"): "BSR Code/ Form 24G Receipt Number",
    ("CD", "challan_serial"): "Bank Challan No/DDO Serial Number",
    ("CD", "challan_date"): "Date of 'Bank Challan / Transfer Voucher' (DDMMYYYY)",
    ("CD", "total_deductee_tax"): ("Total Tax Deposited Amount as per Employee/Deductee annexure "
                                   "(Total Sum of column no.L) for the specified challan"),
    ("CD", "deductee_tax"): ("Total Tax Deducted Amount as per Employee/Deductee annexure "
                             "(Total sum of Column no. J) for the specified challan"),
    ("CD", "minor_head"): "Minor Head of Challan",
    ("CD", "interest_allocation"): "Interest Allocation",
    ("CD", "other_allocation"): "Others Allocation",

    ("DD", "batch_number"): "Batch Number",
    ("DD", "deductee_serial"): "Deductee / Party Detail Record No",
    ("DD", "challan_no"): "Challan reference Number",
    ("DD", "record_status"): "Mode",
    ("DD", "pan"): "Deductee's PAN",
    ("DD", "name"): "Name of the Employee",
    ("DD", "section"): "Section Code",
    ("DD", "deposited"): ("Whether the total amount of tax deducted as mentioned in Column (J) "
                          "has been deposited (Yes/No)"),
    ("DD", "payment_date"): "Date of payment / Credited (dd/mm/yyyy)",
    ("DD", "amount_paid"): "Amount Paid/ Credited",
    ("DD", "tds"): "Total Tax Deducted",
    ("DD", "total_tds"): "Total Tax Deposited",
    ("DD", "deduction_date"): "Date of deduction (dd/mm/yyyy)",
    ("DD", "deposit_date"): "Date of deposit (dd/mm/yyyy)",
    ("DD", "reason_code"): "Reason for non-deduction /lower deduction/higher deduction",
    ("DD", "certificate_no"): ("Certificate number(s) of the certificate under section 395(1) "
                               "issued for non-deduction /lower deduction"),
}


def normalise(value) -> str:
    """The workbook's typography, reduced to the plainest form of the same text."""
    text = "" if value is None else str(value)
    text = (text.replace("\u2018", "'").replace("\u2019", "'")
                .replace("\u201c", '"').replace("\u201d", '"')
                .replace("\u2013", "-").replace("\u2014", "-"))
    return re.sub(r"\s+", " ", text).strip()


def find_workbook() -> Path | None:
    for root in WORKBOOK_CANDIDATES:
        if not root.exists():
            continue
        hits = sorted(root.glob("Form*138*.xls*")) + sorted(root.rglob("Form*138*.xls*"))
        for hit in hits:
            if "138" in hit.name:
                return hit
    return None


def read_workbook_blocks(path: Path) -> dict:
    """{record: {field no: (name, type, size)}} as the workbook declares them."""
    workbook = openpyxl.load_workbook(path, read_only=True)
    sheet = workbook["File Format_R"] if "File Format_R" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    titles = {
        "FH": "File Header Record", "BH": "Batch Header Record",
        "CD": "Challan Detail Record", "DD": "Deductee Detail Record",
    }
    blocks: dict = {}
    for record, needle in titles.items():
        rows: dict = {}
        started = False
        for r in range(1, sheet.max_row + 1):
            first = normalise(sheet.cell(row=r, column=1).value)
            second = normalise(sheet.cell(row=r, column=2).value)
            if not started and needle in second and second.startswith("TDS Statement"):
                started = True
                continue
            if started:
                if re.fullmatch(r"\d+", first):
                    rows[int(first)] = (
                        second,
                        normalise(sheet.cell(row=r, column=4).value),
                        normalise(sheet.cell(row=r, column=5).value),
                    )
                elif rows:
                    break
        blocks[record] = rows
    workbook.close()
    return blocks


WORKBOOK = find_workbook()


class TestLayoutMaps(unittest.TestCase):
    def test_record_widths(self):
        self.assertEqual(layout138.WIDTHS, {"FH": 18, "BH": 72, "CD": 30, "DD": 45})
        import layout140
        self.assertEqual(layout138.WIDTHS, layout140.WIDTHS,
                         "the 138 and 140 record shapes must stay identical")

    def test_file_type_and_label(self):
        self.assertEqual(layout138.FILE_TYPE, "SL1")
        self.assertEqual(layout138.FORM_LABEL, "138")

    def test_no_rate_and_no_140_only_fields(self):
        """The salary form has no rate, no cash-withdrawal amounts and no Form 121."""
        dd = set(layout138.DD_FIELDS.values())
        self.assertNotIn("rate", layout138.DD_FIELDS)
        for number in (22, 23, 28, 31):
            self.assertNotIn(number, dd, f"DD field {number} must stay unmapped on Form 138")

    def test_dd_serial_and_challan_reference_are_not_swapped(self):
        """Field 4 is the record's own serial and field 5 points at the challan.

        That is the reverse of legacy 24Q, where 4 was the challan and 5 the
        record, so this pins the orientation the writer and the linter share.
        """
        self.assertEqual(layout138.DD_FIELDS["deductee_serial"], 4)
        self.assertEqual(layout138.DD_FIELDS["challan_no"], 5)

    def test_field_table_lists_every_mapped_field(self):
        for record, mapping in (("FH", layout138.FH_FIELDS), ("BH", layout138.BH_FIELDS),
                                ("CD", layout138.CD_FIELDS), ("DD", layout138.DD_FIELDS)):
            table = layout138.field_table(record)
            numbers = [row["field"] for row in table]
            self.assertEqual(numbers, sorted(numbers), record)
            self.assertEqual(set(numbers) - {1, 2}, set(mapping.values()) - {1, 2}, record)


class TestFixture138(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.st = Statement.load(FIX / "example_138.json")
        cls.tmp = Path(tempfile.mkdtemp(prefix="tds-test-138-"))
        cls.manifest = build_txt.write(cls.st, cls.tmp)
        cls.txt = cls.tmp / cls.manifest["file"]
        cls.raw = cls.txt.read_bytes()
        cls.lines = cls.raw.decode("ascii").split("\r\n")
        if cls.lines and cls.lines[-1] == "":
            cls.lines.pop()

    def records(self, record_type: str):
        return [line.split("^") for line in self.lines if line.split("^")[1] == record_type]

    def test_record_widths(self):
        widths = {"FH": 18, "BH": 72, "CD": 30, "DD": 45}
        for i, line in enumerate(self.lines, 1):
            parts = line.split("^")
            self.assertEqual(parts[0], str(i), f"line {i} numbering")
            self.assertIn(parts[1], widths, f"line {i} record type")
            self.assertEqual(len(parts), widths[parts[1]], f"line {i} {parts[1]} width")

    def test_head_values(self):
        fh, bh = self.records("FH")[0], self.records("BH")[0]
        self.assertEqual(fh[2], "SL1", "FH field 3 is the file type")
        self.assertEqual(bh[2], "1", "BH field 3 batch number")
        self.assertEqual(bh[3], str(len(self.st.challans)), "BH field 4 counts the challans")
        self.assertEqual(bh[4], "138", "BH field 5 is the form number")

    def test_deductee_serial_is_global_and_strictly_increasing(self):
        serials = [int(parts[3]) for parts in self.records("DD")]
        self.assertEqual(serials, list(range(1, len(self.st.deductees) + 1)))
        self.assertTrue(all(b > a for a, b in zip(serials, serials[1:])))

    def test_each_dd_points_at_the_open_challan(self):
        open_cd = None
        for parts in (line.split("^") for line in self.lines):
            if parts[1] == "CD":
                open_cd = parts[3]
            elif parts[1] == "DD":
                self.assertEqual(parts[4], open_cd, "DD field 5 is the challan reference")

    def test_file_hygiene(self):
        self.assertTrue(self.raw.endswith(b"\r\n"))
        self.assertFalse(self.raw.startswith(b"\xef\xbb\xbf"))
        self.raw.decode("ascii")
        self.assertEqual(self.raw.replace(b"\r\n", b"").count(b"\n"), 0, "no bare LF")
        self.assertEqual(self.lines[-1].split("^")[1], "DD")

    def test_every_populated_field_is_mapped(self):
        """Nothing may land in a field the layout does not own.

        The salary form declares several positions not applicable (a rate, the
        cash-withdrawal amounts, Form 121). If a value ever appears in one of
        them, this fails with the position named.
        """
        maps = {"FH": layout138.FH_FIELDS, "BH": layout138.BH_FIELDS,
                "CD": layout138.CD_FIELDS, "DD": layout138.DD_FIELDS}
        for i, line in enumerate(self.lines, 1):
            parts = line.split("^")
            mapped = set(maps[parts[1]].values())
            stray = [n for n, value in enumerate(parts, 1) if value and n not in mapped]
            self.assertEqual(stray, [], f"line {i} ({parts[1]}) has values in unmapped fields {stray}")

    def test_no_rate_is_written(self):
        for parts in self.records("DD"):
            self.assertEqual(parts[27], "", "DD field 28 ('Remarks 11') must stay empty on Form 138")

    def test_the_two_tax_fields_carry_the_merged_figure(self):
        """Fields 24 and 25 are equal, and both are tax + surcharge + cess.

        The workbook's two remarks require the equality; legacy 24Q spread the
        same money over three fields. Cess is non-zero in the fixture so a
        writer that dropped it would fail here.
        """
        rows = {row.name: row for row in self.st.deductees}
        seen_cess = False
        for parts in self.records("DD"):
            row = rows[parts[8]]
            expected = f"{row.total_tds:.2f}"
            self.assertEqual(parts[23], expected, f"{row.name}: DD field 24")
            self.assertEqual(parts[24], expected, f"{row.name}: DD field 25")
            self.assertEqual(parts[23], parts[24], f"{row.name}: fields 24 and 25 must be equal")
            self.assertEqual(Decimal(expected), row.tds + row.surcharge + row.cess)
            seen_cess = seen_cess or row.cess > 0
        self.assertTrue(seen_cess, "the fixture must exercise a non-zero cess")

    def test_challan_totals_match_the_rows(self):
        """CD fields 21 and 22 are both the sum of the merged per-row tax."""
        per_challan: dict[str, Decimal] = {}
        for parts in self.records("DD"):
            per_challan[parts[4]] = per_challan.get(parts[4], Decimal("0.00")) + Decimal(parts[23])
        for parts in self.records("CD"):
            expected = f"{per_challan.get(parts[3], Decimal('0.00')):.2f}"
            self.assertEqual(parts[20], expected, "CD field 21")
            self.assertEqual(parts[21], expected, "CD field 22")
            self.assertEqual(parts[20], parts[21], "the two annexure sums must agree on this form")

    def test_deposit_date_equals_the_challan_date(self):
        """T-FV-6489: DD field 29 must be the date of the challan it is paid under."""
        open_cd = None
        for parts in (line.split("^") for line in self.lines):
            if parts[1] == "CD":
                open_cd = parts[18]
            elif parts[1] == "DD":
                self.assertEqual(parts[28], open_cd, "DD field 29 must equal CD field 19")

    def test_fixture_is_salary_throughout(self):
        self.assertEqual(self.st.form, "138")
        self.assertEqual([row.section for row in self.st.deductees], ["1002"] * 5)

    def test_rebuild_is_byte_identical(self):
        again = build_txt.write(self.st, tempfile.mkdtemp(prefix="tds-test-138-again-"))
        self.assertEqual(again["sha256"], self.manifest["sha256"])
        self.assertEqual(again["records"], {"FH": 1, "BH": 1, "CD": 2, "DD": 5, "lines": 9})

    def test_manifest_and_short_name(self):
        self.assertEqual(self.manifest["status"], "READY")
        self.assertEqual(self.manifest["file_type"], "SL1")
        self.assertEqual(self.manifest["file"], "EX138Q2.txt")
        self.assertLessEqual(len(build_txt.short_name(self.st)), 12)

    def test_document_round_trips(self):
        again = Statement.from_dict(json.loads(json.dumps(self.st.to_dict())))
        self.assertEqual(self.st.to_dict(), again.to_dict())

    def test_lint_is_clean(self):
        issues = lint.model_issues(self.st) + lint.txt_issues(self.txt, self.st, build_txt.load_layout("138"))
        bad = [str(i) for i in issues if i.level in (lint.ERROR, lint.WARN)]
        self.assertEqual(bad, [], "\n".join(bad))

    def test_quarter_four_is_refused_with_a_reason(self):
        st = Statement.from_dict(json.loads(json.dumps(self.st.to_dict())))
        st.quarter = "Q4"
        with self.assertRaises(build_txt.BuildError) as caught:
            build_txt.write(st, tempfile.mkdtemp(prefix="tds-test-138-q4-"))
        message = str(caught.exception)
        self.assertIn("Q4", message)
        self.assertIn("T-FV-2133", message)

    def test_quarter_four_is_refused_for_any_context(self):
        st = Statement.from_dict(json.loads(json.dumps(self.st.to_dict())))
        st.quarter = "q4"
        with self.assertRaises(build_txt.BuildError):
            layout138.Context(st)


@unittest.skipUnless(WORKBOOK is not None and openpyxl is not None,
                     "no Form 138 file-format workbook on this machine")
class TestAgainstWorkbook(unittest.TestCase):
    """The layout maps, held against the department's own workbook.

    The numbers come from the layout modules; the names come from the workbook
    at those numbers. The names in EXPECTED_NAMES were transcribed from the
    workbook once so that a number which moves between versions is caught.
    """

    @classmethod
    def setUpClass(cls):
        cls.blocks = read_workbook_blocks(WORKBOOK)
        cls.maps = {"FH": layout138.FH_FIELDS, "BH": layout138.BH_FIELDS,
                    "CD": layout138.CD_FIELDS, "DD": layout138.DD_FIELDS}

    def test_record_shapes(self):
        self.assertEqual({rec: len(rows) for rec, rows in self.blocks.items()},
                         {"FH": 18, "BH": 72, "CD": 30, "DD": 45})
        self.assertEqual({rec: len(rows) for rec, rows in self.blocks.items()}, layout138.WIDTHS)

    def test_blocks_are_contiguous(self):
        for record, rows in self.blocks.items():
            self.assertEqual(sorted(rows), list(range(1, len(rows) + 1)),
                             f"{record} field numbers must run 1..n without a gap")

    def test_every_mapped_field_matches_the_official_name(self):
        checked = 0
        for (record, key), expected in EXPECTED_NAMES.items():
            number = self.maps[record][key]
            rows = self.blocks[record]
            self.assertIn(number, rows, f"{record} field {number} ({key}) is missing from the workbook")
            actual = rows[number][0]
            self.assertEqual(actual, expected,
                             f"{record} field {number} ({key}): workbook says {actual!r}")
            checked += 1
        self.assertEqual(checked, len(EXPECTED_NAMES))
        mapped = sum(1 for rec, m in self.maps.items()
                     for key in m if key not in ("line_no", "record_type"))
        self.assertEqual(checked, mapped, "every mapped field needs a name in EXPECTED_NAMES")

    def test_no_rate_field_on_this_form(self):
        self.assertEqual(self.blocks["DD"][28][0], "Remarks 11 (for future use)")

    def test_dd29_type_defect_is_still_in_the_workbook(self):
        """DD 29 is typed DECIMAL/7 but the remark and the FVU require a date.

        The writer follows the remark (and the department's own sample, which
        prints 8-character dates there). If the workbook is ever corrected to a
        DATE, this test tells the next person to drop the note about the defect
        in layout138 and in references/form-138-layout.md.
        """
        name, dtype, size = self.blocks["DD"][29]
        self.assertEqual(name, "Date of deposit (dd/mm/yyyy)")
        self.assertEqual((dtype, size), ("DECIMAL", "7"),
                         "the workbook no longer types DD 29 as DECIMAL/7: update the defect note")


if __name__ == "__main__":
    unittest.main(verbosity=2)
