#!/usr/bin/env python3
"""Books to return, on synthetic Tally-style and Zoho-style ledger exports."""

from __future__ import annotations

import csv
import subprocess
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

import recon  # noqa: E402
from model import Statement  # noqa: E402

FIX = ROOT / "tests" / "fixtures"

# The fixture statement: 1023 = 33500, 1027 = 25000, 1031 = 30000.
TALLY_ROWS = [  # sign-inverted: credits negative, the way Tally often exports
    ["TDS LEDGER EXPORT 01-07-2026 TO 30-09-2026"],
    ["Date", "Particulars", "Ledger", "Vch No", "Amount"],
    ["01-07-2026", "Opening Balance", "TDS Payable - 194C", "", "0"],
    ["31-07-2026", "SAMPLE INFRA PROJECTS", "TDS Payable - 194C", "J1", "-15000"],
    ["31-07-2026", "EXAMPLE TRADERS", "TDS Payable - 194Q", "J2", "-30000"],
    ["31-08-2026", "MEENAKSHI CONSULTING", "TDS on Professional 194J(a)", "J3", "-10000"],
    ["31-08-2026", "DEMO CARRIERS", "TDS Payable - 194C", "J4", "-2000"],
    ["30-09-2026", "MEENAKSHI CONSULTING", "TDS on Professional 194J(a)", "J5", "-14000"],
    ["30-09-2026", "SAMPLE INFRA PROJECTS", "TDS Payable - 194C", "J6", "-16500"],
    ["07-08-2026", "Bank - challan 90001", "TDS Payable - 194C", "P1", "45000"],
    ["30-09-2026", "ROUNDING", "Suspense", "J7", "-5"],
    ["", "Total", "", "", "-42505"],
]


def write_csv(rows) -> Path:
    path = Path(tempfile.mkdtemp()) / "ledger.csv"
    with path.open("w", newline="") as fh:
        csv.writer(fh).writerows(rows)
    return path


class TestTallySignInverted(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.st = Statement.load(FIX / "example_140.json")
        cls.entries = recon.read_ledger(write_csv(TALLY_ROWS))
        cls.result = recon.reconcile(cls.st, cls.entries)
        cls.by = {r["section"]: r for r in cls.result["rows"]}

    def test_negative_amounts_are_read_as_credits(self):
        self.assertEqual(self.result["credits"], Decimal("87505.00"))
        self.assertEqual(self.result["debits"], Decimal("45000.00"))

    def test_sections_come_from_the_ledger_name(self):
        self.assertEqual(self.by["1023"]["books"], Decimal("33500.00"))
        self.assertEqual(self.by["1031"]["books"], Decimal("30000.00"))
        self.assertEqual(self.by["1027"]["books"], Decimal("24000.00"))

    def test_differences_are_listed_row_wise(self):
        diffs = {r["section"]: r["books_less_statement"] for r in self.result["differences"]}
        self.assertEqual(diffs, {"1027": Decimal("-1000.00"), recon.UNMAPPED: Decimal("5.00")})
        self.assertEqual(self.by["1023"]["statement_less_challans"], Decimal("0.00"))
        self.assertEqual([e["particulars"] for e in self.result["unmapped"]], ["ROUNDING"])

    def test_opening_and_total_rows_are_not_entries(self):
        self.assertEqual(len(self.entries), 8)

    def test_positive_credit_sign_flips_the_reading(self):
        res = recon.reconcile(self.st, recon.read_ledger(write_csv(TALLY_ROWS), credit_sign="positive"))
        self.assertEqual(res["credits"], Decimal("45000.00"))


class TestDebitCreditColumns(unittest.TestCase):
    def test_debit_credit_columns_and_drcr_suffix(self):
        rows = [["Date", "Account", "Section", "Debit", "Credit"],
                ["31-07-2026", "TDS Payable", "1023", "", "33,500.00"],
                ["31-07-2026", "TDS Payable", "1031", "", "-30000"],      # sign ignored on a side column
                ["31-08-2026", "TDS Payable", "1027", "", "25000"]]
        st = Statement.load(FIX / "example_140.json")
        res = recon.reconcile(st, recon.read_ledger(write_csv(rows)))
        self.assertEqual(res["differences"], [])
        rows = [["Date", "Ledger", "Amount"],
                ["31-07-2026", "TDS 194C", "33500.00 Cr"], ["31-07-2026", "TDS 194Q", "30000 Cr"],
                ["31-08-2026", "TDS 194J(a)", "25000 Cr"], ["07-08-2026", "TDS 194C", "45000 Dr"]]
        res = recon.reconcile(st, recon.read_ledger(write_csv(rows)))
        self.assertEqual(res["differences"], [])
        self.assertEqual(res["debits"], Decimal("45000.00"))

    def test_expense_ledger_compares_amount_paid(self):
        rows = [["Date", "Ledger", "Debit", "Credit"],
                ["31-07-2026", "Contract charges 194C", "3350000", ""]]
        st = Statement.load(FIX / "example_140.json")
        res = recon.reconcile(st, recon.read_ledger(write_csv(rows)), kind="expense")
        by = {r["section"]: r for r in res["rows"]}
        self.assertEqual(by["1023"]["books_less_statement"], Decimal("0.00"))
        self.assertEqual(by["1023"]["challans"], None)

    def test_cli_writes_the_sheet_and_flags_the_difference(self):
        out = Path(tempfile.mkdtemp()) / "recon.xlsx"
        run = subprocess.run([sys.executable, str(SCRIPTS / "tds.py"), "recon",
                              str(FIX / "example_140.json"), "--ledger", str(write_csv(TALLY_ROWS)),
                              "-o", str(out)], capture_output=True, text=True)
        self.assertEqual(run.returncode, 1, run.stderr)
        self.assertIn("1027", run.stdout)
        from openpyxl import load_workbook
        self.assertIn("Books Reconciliation", load_workbook(out).sheetnames)


if __name__ == "__main__":
    unittest.main(verbosity=2)
