#!/usr/bin/env python3
"""Form 144 layout tests: the numbers, the shared shape, and the extras.

    python3 -m unittest discover -s tests -p 'test_layout144.py' -v

The field numbers checked here were extracted with openpyxl from
``Form144_27Q_Q1-Q4_v1.2_22072026.xlsx``, sheet ``FIleformat_R``, and are pinned
twice: in the literal table below and, where that workbook exists on this
machine, against the workbook itself. The DD block sits in rows 154 to 198, so
field n is workbook row 153 + n.

What these tests prove, and what they do not: they prove that the records
``layout144`` writes have the shape the department's workbook documents, that
the ten Form 144-only figures are wired to ``meta.non_resident`` and are never
invented, and that the fixture builds, lints and round-trips. They cannot prove
the utility accepts the file, and they do not check rules that live only in the
utility - the Form 144-only ones are named in ``references/form-144-layout.md``.
"""

from __future__ import annotations

import json
import re
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

import build_txt  # noqa: E402
import codes  # noqa: E402
import layout140  # noqa: E402
import layout144  # noqa: E402
import lint  # noqa: E402
from model import ModelError, Statement  # noqa: E402

FIX = ROOT / "tests" / "fixtures"
DOC = ROOT / "references" / "form-144-layout.md"

# Where the official workbook may be sitting. The comparison against it skips
# when the workbook is not on this machine, so the suite runs anywhere.
WORKBOOK_DIRS = [
    Path.home() / ".tds-utils" / "format",
    Path.home() / "TDS-Skill-Research" / "official",
    Path.home() / "TDS-Skill-Research" / "downloads",
    Path.home() / "Downloads",
]
WORKBOOK_SHEET = "FIleformat_R"
DD_FIRST_ROW = 154
DD_LAST_ROW = 198

# Field number, key in layout144.DD_FIELDS (None where this toolkit writes
# nothing), the workbook's field name, data type, size and M/O. Extracted from
# the workbook; the labels below are the workbook's own wording, odd spacing
# and all.
DD_TABLE = [
    ( 1, 'line_no'               , 'Line Number', 'INTEGER', '9', 'M'),
    ( 2, 'record_type'           , 'Record Type', 'CHAR', '2', 'M'),
    ( 3, 'batch_number'          , 'Batch Number', 'INTEGER', '9', 'M'),
    ( 4, 'deductee_serial'       , 'Deductee / Party Detail Record No', 'INTEGER', '9', 'M'),
    ( 5, 'challan_no'            , 'Challan reference Number', 'INTEGER', '9', 'M'),
    ( 6, 'record_status'         , 'Mode', 'CHAR', '1', 'M'),
    ( 7, None                    , 'Last deductee PAN ( Used for Verification) (Not applicable)', 'NA', '0', 'O'),
    ( 8, 'pan'                   , "Deductee's PAN", 'CHAR', '10', 'M'),
    ( 9, 'name'                  , 'Name of the Deductee', 'CHAR', '75', 'M'),
    (10, None                    , 'Remarks 1 (for future use)', 'NA', '0', 'O'),
    (11, None                    , 'Remarks 2 (for future use)', 'NA', '0', 'O'),
    (12, 'status'                , 'Status', 'INTEGER', '2', 'M'),
    (13, None                    , 'Remarks 4 (for future use)', 'NA', '0', 'O'),
    (14, 'opting_out'            , 'Whether deductee opting out of taxation regime u/s 202(1)', 'CHAR', '1', 'M'),
    (15, 'section'               , 'Section Code', 'CHAR', '4', 'M'),
    (16, 'nature_of_remittance'  , 'Nature of Remittance', 'INTEGER', '2', 'O'),
    (17, 'country_code'          , 'Country to which remittance is made', 'INTEGER', '3', 'M'),
    (18, 'deposited'             , 'Whether the total amount of tax deducted as mentioned in Column (S) has been deposited (Yes/No)', 'CHAR', '1', 'M'),
    (19, 'payment_date'          , 'Date of payment / Credited (dd/mm/yyyy)', 'DATE', '8', 'M'),
    (20, 'amount_paid'           , 'Amount Paid/ Credited', 'DECIMAL', '15', 'M'),
    (21, 'grossed_up'            , 'Whether the amount paid or credited is grossed up as per section 393(10)', 'CHAR', '1', 'M'),
    (22, 'cash_1cr'              , 'Amount of cash withdrawal in excess of Rs. 1 crore as referred to in section 393(3) [Table: Sl. No. 5] (in case of person other than co- operative society)', 'DECIMAL', '15', 'O'),
    (23, 'cash_3cr'              , 'Amount of cash withdrawal which is in excess of Rs. 3 Crore for cases covered by section 393(3) [Table: Sl. No. 5] (in case of co- operative societies)', 'DECIMAL', '15', 'O'),
    (24, 'tds'                   , 'Total Tax Deducted', 'DECIMAL', '15', 'M'),
    (25, 'total_tds'             , 'Total Tax Deposited', 'DECIMAL', '15', 'M'),
    (26, None                    , 'Last total Tax deducted (Used for Verification) (Not Applicable)', 'NA', '0', 'O'),
    (27, 'deduction_date'        , 'Date of deduction (dd/mm/yyyy)', 'DATE', '8', 'M'),
    (28, 'rate'                  , 'Rate at which Tax Deducted', 'DECIMAL', '7', 'M'),
    (29, None                    , 'Remarks 9 (for future use)', 'NA', '0', 'O'),
    (30, 'rate_basis'            , 'Whether rate of TDS is as per IT Act (a) or DTAA (b)', 'CHAR', '1', 'M'),
    (31, 'form15ca'              , 'Unique acknowledgement number of the corresponding Form No. 15CA [FN 144], if applicable', 'CHAR', '26', 'O'),
    (32, 'reason_code'           , 'Reason for non-deduction /lower deduction/higher deduction (refer Notes 1 to 9)', 'CHAR', '1', 'O'),
    (33, 'certificate_no'        , 'Certificate number(s) of the certificate under section 395(1) / 395(2) issued for non-deduction /lower deduction', 'CHAR', '15', 'O'),
    (34, None                    , 'Filler 1', '', '', ''),
    (35, None                    , 'Filler 2', '', '', ''),
    (36, None                    , 'Filler 3', '', '', ''),
    (37, 'address_abroad'        , 'Address in Country of Residence', 'CHAR', '150', 'O'),
    (38, 'email'                 , 'Email ID', 'CHAR', '75', 'O'),
    (39, 'phone'                 , 'Contact Number (ISD Code,land line or mob)', 'INTEGER', '15', 'O'),
    (40, None                    , 'Filler 7', 'NA', '0', 'O'),
    (41, 'tin'                   , 'Tax Identification Number/Unique Identification Number issued by the country of Residence', 'CHAR', '25', 'O'),
    (42, None                    , 'Last Total Tax Deposited ( Used for Verification) (Not applicable)', 'NA', '0', 'O'),
    (43, None                    , 'Filler 10', 'NA', '0', 'O'),
    (44, None                    , 'Filler 11', 'NA', '0', 'O'),
    (45, None                    , 'Record Hash (Not applicable)', 'NA', '0', 'O'),
]

# The fields whose populated values the fixture pins, line by line. The
# expected dict for each deductee is the whole record: a field that gains a
# value it should not have, or loses one, fails the comparison.
EXPECTED_DD_1 = {
    1: "4", 2: "DD", 3: "1", 4: "1", 5: "1", 6: "O", 8: "AACCZ4321R",
    9: "ZUIDERZEE HOLDINGS B.V.", 12: "1", 14: "N", 15: "1042", 16: "27",
    17: "175", 18: "Y", 19: "31072026", 20: "6325510.00", 21: "N",
    24: "316275.50", 25: "316275.50", 27: "31072026", 28: "5.0000", 30: "A",
    37: "HOOFDWEG 12, 1096 AC AMSTERDAM", 38: "tax@zuiderzee.example",
    39: "31201234567", 41: "NL812345678B01",
}
EXPECTED_DD_2 = {
    1: "5", 2: "DD", 3: "1", 4: "2", 5: "1", 6: "O", 8: "AABCM7654N",
    9: "MERLION TECHNOLOGIES PTE LTD", 12: "1", 14: "N", 15: "1057", 16: "49",
    17: "221", 18: "Y", 19: "31072026", 20: "2500000.00", 21: "N",
    24: "250000.00", 25: "250000.00", 27: "31072026", 28: "10.0000", 30: "B",
    31: "SYNTHETIC15CA0000000000001", 37: "8 MARINA BOULEVARD, #32-01, SINGAPORE 018981",
    38: "accounts@merliontech.example", 39: "6561234567", 41: "200812345K",
}
EXPECTED_DD_3 = {
    1: "7", 2: "DD", 3: "1", 4: "3", 5: "2", 6: "O", 8: "AADCP2468L",
    9: "PRAIRIE CAPITAL LLC", 12: "1", 14: "N", 15: "1044", 16: "27",
    17: "267", 18: "Y", 19: "31082026", 20: "800000.00", 21: "N",
    24: "40000.00", 25: "40000.00", 27: "31082026", 28: "5.0000", 30: "A",
}
EXPECTED_DD_4 = {
    1: "8", 2: "DD", 3: "1", 4: "4", 5: "2", 6: "O", 8: "AAFCT1357M",
    9: "THAMES ADVISORY SERVICES LIMITED", 12: "1", 14: "N", 15: "1057",
    16: "21", 17: "265", 18: "Y", 19: "31082026", 20: "4000000.00", 21: "N",
    24: "300000.00", 25: "300000.00", 27: "31082026", 28: "7.5000", 30: "B",
    32: "A", 33: "TRACESLDC705210", 37: "12 THREADNEEDLE STREET, LONDON EC2R 8AY",
    38: "finance@thamesadvisory.example", 39: "442071234567", 41: "GB123456789",
}

# The fields that carry a rule and so must be explained where the reviewer
# reads, which is FIELD_NOTES. The list is the DD record's; FH, BH and CD carry
# the same rules as on Form 140 and the same notes.
RULE_BEARING_DD_FIELDS = {
    4, 5, 6, 8, 12, 14, 15, 16, 17, 18, 19, 21, 24, 25, 27, 28, 30, 31, 32,
    33, 37, 38, 39, 41,
}


def norm(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def find_workbook() -> Path | None:
    for root in WORKBOOK_DIRS:
        if not root.exists():
            continue
        hits = sorted(root.glob("Form144*.xlsx"))
        if hits:
            return hits[0]
    return None


class TestSharedShape(unittest.TestCase):
    """Form 144 reuses Form 140's records; that reuse has to stay exact."""

    def test_widths(self):
        self.assertEqual(layout144.WIDTHS, {"FH": 18, "BH": 72, "CD": 30, "DD": 45})

    def test_file_type_and_form_number(self):
        self.assertEqual(layout144.FILE_TYPE, "NS1")
        self.assertEqual(layout144.FORM_LABEL, "144")

    def test_fh_bh_and_cd_maps_are_form_140s(self):
        for rec in ("FH", "BH", "CD"):
            self.assertEqual(getattr(layout144, f"{rec}_FIELDS"),
                             getattr(layout140, f"{rec}_FIELDS"), rec)

    def test_the_shared_records_build_identically(self):
        st = Statement.load(FIX / "example_144.json")
        ctx144, ctx140 = layout144.Context(st), layout140.Context(st)
        self.assertEqual(layout144.build_fh(ctx144), layout140.build_fh(ctx140))
        self.assertEqual(
            layout144.build_cd(ctx144, st.challans[0], st.deductees[:2], 1),
            layout140.build_cd(ctx140, st.challans[0], st.deductees[:2], 1))
        bh144, bh140 = layout144.build_bh(ctx144), layout140.build_bh(ctx140)
        field = layout144.BH_FIELDS["form_number"]
        self.assertEqual(bh144.pop(field), "144")
        self.assertEqual(bh140.pop(field), "140")
        self.assertEqual(bh144, bh140)

    def test_every_extra_sits_at_the_field_the_workbook_gives_it(self):
        self.assertEqual(set(layout144.NON_RESIDENT_KEYS),
                         {k for k in layout144.DD_FIELDS if k in layout144.NON_RESIDENT_KEYS})
        for key, no in layout144.NON_RESIDENT_KEYS.items():
            self.assertEqual(layout144.DD_FIELDS[key], no, key)
        for key in layout144.NON_RESIDENT_MANDATORY:
            self.assertIn(key, layout144.NON_RESIDENT_KEYS, key)

    def test_the_rule_bearing_fields_carry_a_note(self):
        noted = {r["field"] for rec in ("FH", "BH", "CD", "DD")
                 for r in layout144.field_table(rec)}
        missing = RULE_BEARING_DD_FIELDS - noted
        self.assertEqual(missing, set(), f"DD fields with a rule and no note: {sorted(missing)}")

    def test_flags_and_rates_are_written_the_way_the_format_wants_them(self):
        st = Statement.load(FIX / "example_144.json")
        block = st.meta["non_resident"]["AACCZ4321R"]
        block["opting_out"] = " n "
        block["grossed_up"] = "y"
        block["rate_basis"] = "b"
        rec = layout144.build_dd(layout144.Context(st), st.deductees[0], 1, 1)
        self.assertEqual(rec[layout144.DD_FIELDS["opting_out"]], "N")
        self.assertEqual(rec[layout144.DD_FIELDS["grossed_up"]], "Y")
        self.assertEqual(rec[layout144.DD_FIELDS["rate_basis"]], "B")
        self.assertEqual(layout144._rate("7.5"), "7.5000")
        self.assertEqual(layout144._rate(""), "")


class TestFixture(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.st = Statement.load(FIX / "example_144.json")
        cls.out = Path(tempfile.mkdtemp(prefix="tds-144-"))
        cls.manifest = build_txt.write(cls.st, cls.out)
        cls.txt = cls.out / cls.manifest["file"]
        cls.raw = cls.txt.read_bytes()
        cls.lines = cls.raw.decode("ascii").split("\r\n")
        if cls.lines and cls.lines[-1] == "":
            cls.lines.pop()

    # ---- the file -------------------------------------------------------
    def test_record_widths(self):
        widths = {"FH": 18, "BH": 72, "CD": 30, "DD": 45}
        for i, line in enumerate(self.lines, 1):
            parts = line.split("^")
            self.assertEqual(parts[0], str(i), f"line {i} numbering")
            self.assertIn(parts[1], widths, f"line {i} record type")
            self.assertEqual(len(parts), widths[parts[1]], f"line {i} {parts[1]} width")

    def test_fh_carries_the_file_type_and_bh_the_form_number(self):
        fh, bh = (l.split("^") for l in self.lines[:2])
        self.assertEqual(fh[layout144.FH_FIELDS["file_type"] - 1], "NS1")
        self.assertEqual(fh[layout144.FH_FIELDS["record_type"] - 1], "FH")
        self.assertEqual(bh[layout144.BH_FIELDS["form_number"] - 1], "144")

    def test_deductee_serial_is_global_and_increases(self):
        serials = [int(l.split("^")[layout144.DD_FIELDS["deductee_serial"] - 1])
                   for l in self.lines if l.split("^")[1] == "DD"]
        self.assertEqual(serials, sorted(set(serials)))
        self.assertEqual(serials, list(range(1, len(self.st.deductees) + 1)))

    def test_file_hygiene(self):
        self.assertTrue(self.raw.endswith(b"\r\n"))
        self.assertFalse(self.raw.startswith(b"\xef\xbb\xbf"))
        self.raw.decode("ascii")  # raises if any byte is outside ASCII
        self.assertNotIn(b"\n", self.raw.replace(b"\r\n", b""))

    def test_the_manifest_names_the_fixture(self):
        self.assertEqual(self.manifest["file"], "EX144Q2.txt")
        self.assertEqual(self.manifest["status"], "READY")
        self.assertEqual({k: v for k, v in self.manifest["records"].items() if k != "lines"},
                         {"FH": 1, "BH": 1, "CD": 2, "DD": 4})

    # ---- what is in the records ----------------------------------------
    def test_every_populated_field_is_in_the_module_map(self):
        ctx = layout144.Context(self.st)
        records = {
            "FH": (layout144.build_fh(ctx), layout144.FH_FIELDS),
            "BH": (layout144.build_bh(ctx), layout144.BH_FIELDS),
            "CD": (layout144.build_cd(ctx, self.st.challans[0], self.st.deductees[:2], 1),
                   layout144.CD_FIELDS),
            "DD": (layout144.build_dd(ctx, self.st.deductees[0], 1, 1), layout144.DD_FIELDS),
        }
        for rec, (built, mapping) in records.items():
            written, mapped = set(built), set(mapping.values())
            self.assertTrue(written <= mapped,
                            f"{rec}: written but not in the map: {sorted(written - mapped)}")
            self.assertEqual(mapped - written, {mapping["line_no"]},
                             f"{rec}: in the map but never written")
            self.assertLessEqual(max(mapped), layout144.WIDTHS[rec], rec)

    def test_the_deductee_rows_are_exactly_what_is_expected(self):
        rows = [{i: v for i, v in enumerate(l.split("^"), 1) if v}
                for l in self.lines if l.split("^")[1] == "DD"]
        self.assertEqual(len(rows), 4)
        for got, expected in zip(rows, (EXPECTED_DD_1, EXPECTED_DD_2, EXPECTED_DD_3, EXPECTED_DD_4)):
            self.assertEqual(got, expected)

    def test_the_form_144_extras_are_written_and_left_blank_when_absent(self):
        rows = [l.split("^") for l in self.lines if l.split("^")[1] == "DD"]
        keys = layout144.NON_RESIDENT_KEYS
        blocks = self.st.meta["non_resident"]
        for row, deductee in zip(rows, self.st.deductees):
            block = blocks.get(deductee.pan, {})
            for key, no in keys.items():
                expected = norm(block.get(key, ""))
                if key in ("opting_out", "grossed_up", "rate_basis") and expected:
                    expected = expected.upper()[:1]
                self.assertEqual(row[no - 1], expected, f"{deductee.pan} {key}")

    def test_the_challan_record_is_whole_rupees_and_the_rows_are_not(self):
        cds = [l.split("^") for l in self.lines if l.split("^")[1] == "CD"]
        self.assertEqual([cd[7] for cd in cds], ["566276.00", "340000.00"])
        self.assertEqual([cd[11] for cd in cds], ["566276.00", "343400.00"])
        self.assertEqual(cds[1][8], "3400.00")           # interest on the late challan
        self.assertEqual(cds[1][23], "3400.00")           # and the interest allocation
        self.assertEqual(cds[0][24], "0.00")              # no penalty anywhere
        dd = [l.split("^") for l in self.lines if l.split("^")[1] == "DD"]
        self.assertEqual(dd[0][23], "316275.50")          # paise are allowed here
        self.assertEqual(cds[0][20], "566275.50")         # CD 21, the rows' deposit

    def test_the_challans_carry_their_rows_and_no_more(self):
        for row in self.st.allocation():
            self.assertGreaterEqual(row["remaining"], Decimal("0.00"), row)
        first = self.st.allocation()[0]
        self.assertEqual(first["unallocated"], Decimal("0.50"))

    def test_round_trip_and_rebuild(self):
        again = Statement.from_dict(json.loads(json.dumps(self.st.to_dict())))
        self.assertEqual(again.to_dict(), self.st.to_dict())
        second = build_txt.write(again, tempfile.mkdtemp(prefix="tds-144-again-"))
        self.assertEqual(second["sha256"], self.manifest["sha256"])
        self.assertEqual(second["bytes"], len(self.raw))

    # ---- what happens when the document does not say ---------------------
    def test_a_row_without_the_extras_is_written_blank_not_guessed(self):
        st = Statement.load(FIX / "example_144.json")
        st.meta.pop("non_resident")
        out = Path(tempfile.mkdtemp(prefix="tds-144-bare-"))
        manifest = build_txt.write(st, out)
        lines = (out / manifest["file"]).read_bytes().decode("ascii").split("\r\n")
        first = next(l.split("^") for l in lines if l.split("^")[1] == "DD")
        for no in layout144.NON_RESIDENT_KEYS.values():
            self.assertEqual(first[no - 1], "", f"DD field {no} should be blank")

    def test_a_malformed_non_resident_block_is_refused(self):
        for bad in (["AACCZ4321R"],
                    {"AACCZ4321R": "1"},
                    {"AACCZ4321R": {"staus": "1"}}):
            st = Statement.load(FIX / "example_144.json")
            st.meta["non_resident"] = bad
            with self.assertRaises(ModelError, msg=repr(bad)):
                build_txt.render(st, layout144)

    # ---- lint -----------------------------------------------------------
    def test_lint_is_clean_against_form_144s_own_section_table(self):
        """Form 144's Annexure 2 is checked from assets/annexures_144.json."""
        issues = lint.model_issues(self.st) + lint.txt_issues(self.txt, self.st, layout144)
        # The fixture's second challan is paid late on purpose; the interest
        # working flags that, which is not a layout matter.
        bad = [str(i) for i in issues if i.level in (lint.ERROR, lint.WARN)
               and i.code != "INTEREST_SHORT"]
        self.assertEqual(bad, [], "\n".join(bad))

    def test_a_form_140_only_code_is_refused_on_form_144(self):
        st = Statement.load(FIX / "example_144.json")
        st.deductees[0].section = "1023"
        codes_ = {i.code for i in lint.model_issues(st) if i.level == lint.ERROR}
        self.assertIn("SECTION", codes_)


class TestAgainstTheWorkbook(unittest.TestCase):
    """The pinned numbers, read back out of the department's own workbook."""

    @classmethod
    def setUpClass(cls):
        cls.path = find_workbook()
        if cls.path is None:
            raise unittest.SkipTest(
                "Form144*.xlsx is not on this machine; run with the workbook to "
                "check the pinned numbers against it")
        import openpyxl
        wb = openpyxl.load_workbook(cls.path, data_only=True, read_only=True)
        cls.ws = wb[WORKBOOK_SHEET]

    def cell(self, row, column):
        return norm(self.ws.cell(row=row, column=column).value)

    def test_the_dd_table_is_the_workbooks(self):
        for field, _key, label, dtype, size, mo in DD_TABLE:
            row = DD_FIRST_ROW - 1 + field
            self.assertEqual(self.cell(row, 1), str(field), f"row {row} field number")
            self.assertEqual(self.cell(row, 2), norm(label), f"DD {field} name")
            self.assertEqual(self.cell(row, 4), dtype, f"DD {field} type")
            self.assertEqual(self.cell(row, 5), size, f"DD {field} size")
            self.assertEqual(self.cell(row, 6), mo, f"DD {field} M/O")

    def test_the_module_map_is_the_dd_table(self):
        from_table = {row[0]: row[1] for row in DD_TABLE}
        in_module = {no: name for name, no in layout144.DD_FIELDS.items()}
        for field, key in in_module.items():
            self.assertEqual(from_table.get(field), key, f"DD {field}")
        for field, key in from_table.items():
            if key is None:
                self.assertNotIn(field, in_module, f"DD {field} is a filler")

    def test_the_blocks_hold_the_field_counts_this_module_declares(self):
        blocks = {"FH": (22, 39, 18), "BH": (44, 115, 72), "CD": (120, 149, 30),
                  "DD": (DD_FIRST_ROW, DD_LAST_ROW, 45)}
        for rec, (first, last, count) in blocks.items():
            numbers = [self.cell(r, 1) for r in range(first, last + 1)]
            self.assertEqual(numbers, [str(n) for n in range(1, count + 1)], rec)
            self.assertEqual(layout144.WIDTHS[rec], count, rec)

    def test_form_144_has_its_own_section_table(self):
        codes_144 = {self.cell(r, 4) for r in range(242, 272)}
        self.assertEqual(len(codes_144), 30)
        self.assertIn("1004", codes_144)
        self.assertEqual(sorted(codes_144 & set(codes.SECTIONS)),
                         ["1004", "1058", "1059", "1060", "1061", "1062", "1063",
                          "1064", "1065", "1066", "1067"])
        self.assertEqual(codes_144, set(codes.SECTIONS_144),
                         "assets/annexures_144.json must be the workbook's Annexure 2; "
                         "rerun extract_annexures.py --form144")

    def test_the_fixtures_codes_are_real_codes(self):
        countries = {self.cell(r, 3): self.cell(r, 2) for r in range(564, 849)}
        statuses = {self.cell(r, 3) for r in range(477, 487)}
        natures = {self.cell(r, 3) for r in range(490, 560)}
        sections_for_nature: dict[str, set] = {}
        for r in range(490, 560):
            nature, section = self.cell(r, 3), self.cell(r, 4)
            if re.fullmatch(r"\d{2}", nature) and re.fullmatch(r"\d{4}", section):
                sections_for_nature.setdefault(nature, set()).add(section)
        st = Statement.load(FIX / "example_144.json")
        for pan, block in st.meta["non_resident"].items():
            if "country_code" in block:
                self.assertIn(block["country_code"], countries, pan)
            if "status" in block:
                self.assertIn(block["status"], statuses, pan)
            if "nature_of_remittance" in block:
                self.assertIn(block["nature_of_remittance"], natures, pan)
        for row in st.deductees:
            block = st.meta["non_resident"][row.pan]
            nature = block.get("nature_of_remittance")
            allowed = sections_for_nature.get(nature, set())
            if allowed:
                self.assertIn(row.section, allowed,
                              f"{row.section} is not one of the sections Annexure 9 "
                              f"allows for nature {nature}")
        self.assertEqual(countries["175"], "NETHERLANDS")
        self.assertEqual(countries["265"], "UNITED KINGDOM")
        self.assertEqual(countries["267"], "UNITED STATES OF AMERICA")

    def test_the_fixtures_reason_flag_is_offered_for_its_section(self):
        flags = {(self.cell(r, 2), self.cell(r, 4)) for r in range(395, 468)}
        st = Statement.load(FIX / "example_144.json")
        for row in st.deductees:
            if row.reason_code:
                self.assertIn((row.reason_code, row.section), flags,
                              f"flag {row.reason_code} is not offered for {row.section}")


class TestReferenceDoc(unittest.TestCase):
    """The reference page is part of the deliverable, so it is checked too."""

    @classmethod
    def setUpClass(cls):
        if not DOC.exists():
            raise unittest.SkipTest(f"{DOC} is missing")
        cls.text = DOC.read_text(encoding="utf-8")

    def _table(self, record: str) -> list[dict]:
        head = f"## {record} record"
        start = self.text.index(head)
        rest = self.text[start + len(head):]
        end = rest.find("\n## ")
        block = rest if end < 0 else rest[:end]
        rows = []
        for line in block.splitlines():
            if not line.startswith("| "):
                continue
            cells = [c.strip() for c in line.strip("|").split("|")]
            if not cells or not cells[0].isdigit():
                continue
            rows.append({"field": int(cells[0]), "name": cells[1],
                         "module": cells[6], "remark": cells[7] if len(cells) > 7 else ""})
        return rows

    def test_the_doc_lists_every_field_the_module_writes(self):
        maps = {"FH": layout144.FH_FIELDS, "BH": layout144.BH_FIELDS,
                "CD": layout144.CD_FIELDS, "DD": layout144.DD_FIELDS}
        for rec, mapping in maps.items():
            rows = self._table(rec)
            self.assertEqual(len(rows), layout144.WIDTHS[rec], rec)
            self.assertEqual([r["field"] for r in rows],
                             list(range(1, layout144.WIDTHS[rec] + 1)), rec)
            for r in rows:
                expected = next((k for k, no in mapping.items() if no == r["field"]), None)
                if expected is None:
                    self.assertIn(r["module"], ("-", "\u2013", ""),
                                  f"{rec} {r['field']} should not name a module key")
                else:
                    self.assertEqual(r["module"], f"`{expected}`",
                                     f"{rec} {r['field']}")

    def test_the_doc_names_its_source_and_the_row_ranges(self):
        self.assertIn("Form144_27Q_Q1-Q4_v1.2_22072026.xlsx", self.text)
        self.assertIn("FIleformat_R", self.text)
        for phrase in ("22-39", "44-115", "120-149", "154-198"):
            self.assertIn(phrase, self.text, phrase)

    def test_the_doc_records_the_known_gap_and_the_unknowns(self):
        for phrase in ("1039", "1067", "T-FV-6425", "T-FV-6429", "T-FV-4093",
                       "meta.non_resident", "202(1)"):
            self.assertIn(phrase, self.text, phrase)


if __name__ == "__main__":
    unittest.main(verbosity=2)
