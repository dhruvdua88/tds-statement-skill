#!/usr/bin/env python3
"""Regenerate the reference tables from the department's own file-format workbook.

Every code in this toolkit (state codes, section codes, deductor categories,
reason flags, minor heads) comes from the official spreadsheet, never from
memory. A state code recalled from memory is exactly the kind of error that
produces a statement which looks right and is rejected at upload, so the tables
are generated and the generator is checked.

    python3 scripts/extract_annexures.py                 # find the workbook
    python3 scripts/extract_annexures.py --workbook <xlsx>

Outputs, all committed so users do not need the workbook:

    assets/annexures.json     the code tables the linter reads
    assets/fvu_codes.json     every validation code with its message string
    references/form-140-layout.md   the field tables, with the official remarks
    references/section-codes.md     the code tables, readable

The row map below was verified against the Form 140 file-format workbook
published 22-Jul-2026 (labelled version 1.2). Sheet names and row offsets do
change between editions, so the script asserts the expected shape (18/72/30/45
fields, 36/40/58/16/184/3 annexure rows) and records any mismatch in the output
rather than writing a quietly wrong table.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
ASSETS = ROOT / "assets"
REFERENCES = ROOT / "references"

SHEET = "File Format_R"

BLOCKS = {
    "FH": (22, 39),
    "BH": (44, 115),
    "CD": (120, 149),
    "DD": (154, 198),
}
ANNEXURES = {
    "state_codes": (203, 239),
    "sections": (242, 282),
    "ministry_codes": (286, 344),
    "deductor_categories": (347, 363),
    "state_codes_repeat": (366, 402),
    "reason_flags": (405, 589),
    "minor_heads": (593, 596),
}
EXPECTED = {
    "fields": {"FH": 18, "BH": 72, "CD": 30, "DD": 45},
    "annexures": {"state_codes": 36, "sections": 40, "ministry_codes": 58,
                  "deductor_categories": 16, "reason_flags": 184, "minor_heads": 3},
}

CANDIDATES = [
    Path.home() / ".tds-utils" / "format",
    Path.home() / "TDS-Skill-Research" / "downloads",
    Path.home() / "TDS-Skill-Research" / "src",
    Path.home() / "AWR-TDS-Return-Q2-FY2627" / "research" / "downloads",
    Path.home() / "Downloads",
]

CODE_JSON_CANDIDATES = [
    Path.home() / ".tds-utils" / "fvu" / "fvu_codes.json",
    Path.home() / "AWR-TDS-Return-Q2-FY2627" / "research" / "fvu_run" / "fvu12_codes_extracted.json",
]


def find_workbook(explicit: str | None) -> Path | None:
    if explicit:
        p = Path(explicit).expanduser()
        return p if p.exists() else None
    for root in CANDIDATES:
        if not root.exists():
            continue
        hits = sorted(root.glob("Form140*FileFormat*.xls*")) + sorted(root.glob("Form140*.xlsx"))
        for h in hits:
            if "FileFormat" in h.name or "File Format" in h.name:
                return h
    return None


def find_codes(explicit: str | None) -> Path | None:
    if explicit:
        p = Path(explicit).expanduser()
        return p if p.exists() else None
    for p in CODE_JSON_CANDIDATES:
        if p.exists():
            return p
    return None


def txt(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"[ \t]+", " ", s).strip()


def md(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ").strip()


def extract_fields(ws, r0: int, r1: int) -> list[dict]:
    out = []
    for r in range(r0, r1 + 1):
        sr = txt(ws.cell(row=r, column=1).value)
        if not re.fullmatch(r"\d+", sr):
            continue
        out.append({
            "field": int(sr),
            "name": txt(ws.cell(row=r, column=2).value),
            "form_col": txt(ws.cell(row=r, column=3).value),
            "dtype": txt(ws.cell(row=r, column=4).value),
            "size": txt(ws.cell(row=r, column=5).value),
            "mandatory": txt(ws.cell(row=r, column=6).value),
            "remark": txt(ws.cell(row=r, column=7).value),
        })
    return out


def extract_rows(ws, r0: int, r1: int, cols=(2, 3, 4, 5)) -> list[list[str]]:
    out = []
    for r in range(r0, r1 + 1):
        vals = [txt(ws.cell(row=r, column=c).value) for c in cols]
        if any(vals):
            out.append(vals)
    return out


def build_tables(ws) -> tuple[dict, list[str]]:
    warnings: list[str] = []
    record_fields = {k: extract_fields(ws, *v) for k, v in BLOCKS.items()}
    for rec, rows in record_fields.items():
        if len(rows) != EXPECTED["fields"][rec]:
            warnings.append(
                f"{rec}: {len(rows)} fields found, expected {EXPECTED['fields'][rec]} "
                f"- the row map in this script needs updating for the new edition"
            )

    raw = {k: extract_rows(ws, *v) for k, v in ANNEXURES.items()}
    for key, expected in EXPECTED["annexures"].items():
        # the spreadsheet includes its own header row in each range
        got = len(raw[key]) - 1
        if got != expected:
            warnings.append(f"{key}: {got} rows found, expected {expected} - check the range")

    def drop_header(rows: list[list[str]], header_word: str) -> list[list[str]]:
        out = []
        for row in rows:
            if row and header_word.lower() in row[0].lower() and len(row[0]) < 40 and not row[1]:
                continue
            if row and row[0].lower() in ("state name", "nature of payment", "ministry name",
                                          "particulars", "flag", "category description"):
                continue
            out.append(row)
        return out

    states = {}
    for row in drop_header(raw["state_codes"], "state name"):
        name, code = row[0], row[1]
        if re.fullmatch(r"\d{1,2}", code):
            states[code.zfill(2)] = name

    sections = {}
    for row in drop_header(raw["sections"], "nature of payment"):
        desc, section_ref, code = row[0], row[1], row[2]
        if re.fullmatch(r"\d{4}", code):
            sections[code] = {"description": desc, "section": section_ref}

    categories = {}
    for row in drop_header(raw["deductor_categories"], "category description"):
        desc, value = row[0], row[1]
        if re.fullmatch(r"[A-Za-z]", value):
            categories[value.upper()] = desc

    minor_heads = {}
    for row in drop_header(raw["minor_heads"], "particulars"):
        desc, code = row[0], row[1]
        if re.fullmatch(r"\d{3}", code):
            minor_heads[code] = desc

    reason_flags: dict[str, dict] = {}
    for row in drop_header(raw["reason_flags"], "flag"):
        flag, section_ref, code, meaning = (row + ["", "", "", ""])[:4]
        if not flag or not re.fullmatch(r"[A-Za-z0-9]{1,2}", flag):
            continue
        entry = reason_flags.setdefault(flag.upper(), {"meaning": "", "available_for": []})
        if meaning and not entry["meaning"]:
            entry["meaning"] = meaning
        if code:
            entry["available_for"].append({"section": section_ref, "code": code})

    tables = {
        "state_codes": states,
        "sections": sections,
        "ministry_codes": {row[1].zfill(2): row[0] for row in drop_header(raw["ministry_codes"], "ministry name")
                           if len(row) > 1 and re.fullmatch(r"\d{1,2}", row[1])},
        "deductor_categories": categories,
        "minor_heads": minor_heads,
        "reason_codes": {flag: entry["meaning"] for flag, entry in reason_flags.items()},
        "reason_code_detail": reason_flags,
        "records": record_fields,
        "warnings": warnings,
    }
    return tables, warnings


def write_references(tables: dict, source_name: str) -> None:
    REFERENCES.mkdir(parents=True, exist_ok=True)

    lines = [
        "# Form 140 (resident non-salary) record layout",
        "",
        f"Generated from the department's own file-format workbook, `{source_name}`, "
        "sheet `File Format_R`. Do not hand-edit; rerun `scripts/extract_annexures.py`.",
        "",
        "The statements in this pipeline write these records exactly as described here. "
        "Where a remark imposes a rule, the rule is quoted in full.",
        "",
        "| Record | Fields | Meaning |",
        "|---|---:|---|",
        "| FH | 18 | File header, once per file |",
        "| BH | 72 | Batch header, once per batch |",
        "| CD | 30 | Challan detail, once per challan |",
        "| DD | 45 | Deductee detail, once per deductee |",
        "",
        "A regular statement carries exactly one batch, so a file is: one FH, one BH, "
        "then for each challan a CD immediately followed by its DD records.",
        "",
    ]
    for rec, rows in tables["records"].items():
        lines.append(f"## {rec} record")
        lines.append("")
        lines.append("| # | Field | Form col. | Type | Size | M/O | Remark |")
        lines.append("|---:|---|---|---|---:|---|---|")
        for r in rows:
            lines.append(
                f"| {r['field']} | {md(r['name'])} | {md(r['form_col']) or '–'} | "
                f"{md(r['dtype']) or 'NA'} | {md(r['size']) or '0'} | "
                f"{md(r['mandatory']) or 'NA'} | {md(r['remark'])} |"
            )
        lines.append("")
    (REFERENCES / "form-140-layout.md").write_text("\n".join(lines), encoding="utf-8")

    lines = [
        "# Section codes and other reference tables",
        "",
        f"Generated from `{source_name}`, sheet `File Format_R`. "
        "Do not hand-edit; rerun `scripts/extract_annexures.py`.",
        "",
        "## Annexure 1 - state and union territory codes (BH field 26)",
        "",
        "| Code | State / UT |",
        "|---|---|",
    ]
    for code, name in sorted(tables["state_codes"].items(), key=lambda kv: int(kv[0])):
        lines.append(f"| {code} | {name} |")
    lines += ["", "## Annexure 2 - nature of payment, section and return code (DD field 15)", "",
              "| Return code | Nature of payment | Section (Income-tax Act, 2025) |", "|---|---|---|"]
    for code, meta in sorted(tables["sections"].items()):
        lines.append(f"| {code} | {md(meta['description'])} | {md(meta['section'])} |")
    lines += ["", "## Annexure 4 - deductor and collector category (BH field 32)", "",
              "| Value | Category |", "|---|---|"]
    for value, desc in sorted(tables["deductor_categories"].items()):
        lines.append(f"| {value} | {md(desc)} |")
    lines += ["", "## Annexure 6 - reason flags (DD field 32)", "",
              "These are the codes for a lower, nil or higher deduction, a threshold case, "
              "a transporter and so on. The meanings below come from the workbook; the "
              "detail sheet records which return code each flag is available for.", "",
              "| Flag | Meaning |", "|---|---|"]
    for flag, entry in sorted(tables["reason_code_detail"].items()):
        lines.append(f"| {flag} | {md(entry['meaning'])} |")
    lines += ["", "## Annexure 7 - minor head code (CD field 23)", "",
              "| Code | Particulars |", "|---|---|"]
    for code, desc in sorted(tables["minor_heads"].items()):
        lines.append(f"| {code} | {md(desc)} |")
    lines += ["", "## Annexure 3 - ministry codes (BH field 57)", "",
              "Government deductors only. Code `99` is 'Others' and switches on the "
              "free-text ministry name in BH field 58.", "",
              "| Code | Ministry |", "|---|---|"]
    for code, desc in sorted(tables["ministry_codes"].items()):
        lines.append(f"| {code} | {md(desc)} |")
    lines.append("")
    (REFERENCES / "section-codes.md").write_text("\n".join(lines), encoding="utf-8")


# ---------------------------------------------------------------------------
# Form 144. Its workbook has the same four record blocks at the same rows as
# Form 140's, but its own Annexure 2 (code 1004 and codes 1039 to 1067), so
# lint needs a second section table. Only that table is written as an asset;
# the field tables go to references/form-144-layout.md.
# ---------------------------------------------------------------------------
SHEET_144 = "FIleformat_R"
SECTIONS_144_ROWS = (242, 271)
EXPECTED_SECTIONS_144 = 30

FORM144_PREAMBLE = """# Form 144 (payments to non-residents) record layout

Generated by `scripts/extract_annexures.py --form144` from the department's own
file-format workbook, `Form144_27Q_Q1-Q4_v1.2_22072026.xlsx` (published on the
Protean e-TDS download page as "Form Number 144-27Q - Q1 to Q4_22072026.xlsx"),
sheet `FIleformat_R`. The record blocks sit in rows 22-39 (FH), 44-115 (BH),
120-149 (CD) and 154-198 (DD). The workbook is the authority; regenerate this
page when the department republishes it.

The column *Module key* names the key in `scripts/layout144.py` that writes the
field; a dash means the builder leaves it blank.

## How this form differs from Form 140

- FH, BH and CD are field for field the same as Form 140's. Only the DD record
  differs.
- Ten DD fields that are fillers on Form 140 carry values here: status (12),
  opting out under section 202(1) (14), nature of remittance (16), country (17),
  grossing up (21), rate basis Act or DTAA (30), and the deductee's address,
  e-mail, phone and tax identification number abroad (37, 38, 39, 41). The
  statement document has no field for them, so they travel in
  `meta.non_resident`, keyed by the deductee's PAN, and a value the document does
  not carry is written blank rather than guessed. The utility enforces the
  mandatory ones itself: T-FV-6425 (country), T-FV-6429 (grossing up) and
  T-FV-4093 (rate, which Form 144 must carry).
- Form 144 has its own Annexure 2: code 1004 and codes 1039 to 1067. It is in
  `assets/annexures_144.json`, and lint checks a 144 statement against it.
- What section 202(1) covers could not be established from the material
  available here; DD field 14 is written from the document as given.
"""


def write_form144(path: Path) -> int:
    import openpyxl

    sys.path.insert(0, str(HERE))
    import layout144

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET_144 not in wb.sheetnames:
        print(f"sheet {SHEET_144!r} not found; sheets are {wb.sheetnames}", file=sys.stderr)
        return 2
    ws = wb[SHEET_144]

    sections = {}
    for code, nature, section in ((r[2], r[0], r[1]) for r in extract_rows(ws, *SECTIONS_144_ROWS)):
        if re.fullmatch(r"\d{4}", code):
            sections[code] = {"description": nature, "section": section}
    if len(sections) != EXPECTED_SECTIONS_144:
        print(f"expected {EXPECTED_SECTIONS_144} Annexure 2 rows, found {len(sections)}; "
              "the workbook edition has changed, so nothing was written", file=sys.stderr)
        return 2
    asset = {"sections": sections, "source": {"file": path.name, "sheet": SHEET_144}}
    (ASSETS / "annexures_144.json").write_text(
        json.dumps(asset, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    maps = {"FH": layout144.FH_FIELDS, "BH": layout144.BH_FIELDS,
            "CD": layout144.CD_FIELDS, "DD": layout144.DD_FIELDS}
    lines = [FORM144_PREAMBLE]
    for rec, (r0, r1) in BLOCKS.items():
        keys = {no: key for key, no in maps[rec].items()}
        lines += [f"## {rec} record", "",
                  "| # | Field | Form col. | Type | Size | M/O | Module key | Remark |",
                  "|---:|---|---|---|---:|---|---|---|"]
        for f in extract_fields(ws, r0, r1):
            key = keys.get(f["field"])
            lines.append(f"| {f['field']} | {md(f['name'])} | {md(f['form_col'])} | {md(f['dtype'])} | "
                         f"{md(f['size'])} | {md(f['mandatory'])} | {f'`{key}`' if key else '-'} | "
                         f"{md(f['remark'])} |")
        lines.append("")
    (REFERENCES / "form-144-layout.md").write_text("\n".join(lines), encoding="utf-8")
    print(f"section codes     {len(sections)}")
    print(f"wrote {ASSETS / 'annexures_144.json'}")
    print(f"wrote {REFERENCES / 'form-144-layout.md'}")
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", default=None, help="the official file-format workbook")
    ap.add_argument("--codes", default=None, help="JSON of FVU validation codes, if kept locally")
    ap.add_argument("--form144", default=None,
                    help="the Form 144 workbook: write assets/annexures_144.json and "
                         "references/form-144-layout.md, and nothing else")
    args = ap.parse_args(argv)
    if args.form144:
        return write_form144(Path(args.form144).expanduser())

    wb_path = find_workbook(args.workbook)
    if not wb_path:
        print("The official file-format workbook was not found.\n", file=sys.stderr)
        print("Download it from the department's TDS page (see references/official-downloads.json),", file=sys.stderr)
        print("then run:  python3 scripts/extract_annexures.py --workbook /path/to/Form140_FileFormat.xlsx", file=sys.stderr)
        print("\nThe committed assets/annexures.json is still usable; this only matters when a", file=sys.stderr)
        print("new format edition is published.", file=sys.stderr)
        return 2

    import openpyxl

    wb = openpyxl.load_workbook(wb_path, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        print(f"sheet {SHEET!r} not found; sheets are {wb.sheetnames}", file=sys.stderr)
        return 2
    tables, warnings = build_tables(wb[SHEET])
    tables["generated_from"] = wb_path.name
    tables["sheet"] = SHEET
    assets = {"state_codes": tables["state_codes"], "sections": {
        k: {"description": v["description"], "section": v["section"]} for k, v in tables["sections"].items()},
        "deductor_categories": tables["deductor_categories"],
        "minor_heads": tables["minor_heads"],
        "reason_codes": tables["reason_codes"],
        "annexures": {"ministry_codes": tables["ministry_codes"],
                      "reason_code_detail": tables["reason_code_detail"]},
        "source": {"file": wb_path.name, "sheet": SHEET},
        "warnings": warnings,
    }
    ASSETS.mkdir(parents=True, exist_ok=True)
    (ASSETS / "annexures.json").write_text(json.dumps(assets, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (ASSETS / "form-140-records.json").write_text(
        json.dumps(tables["records"], indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    write_references(tables, wb_path.name)

    codes_path = find_codes(args.codes)
    if codes_path:
        raw = json.loads(codes_path.read_text(encoding="utf-8"))
        codes = {}
        for code, msg in raw.items():
            if isinstance(msg, list):
                msg = msg[0] if msg else ""
            codes[str(code).replace("_", "-")] = str(msg)
        (ASSETS / "fvu_codes.json").write_text(json.dumps(codes, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"validation codes: {len(codes)}")
    else:
        print("validation codes: not found locally, assets/fvu_codes.json left as is")

    print(f"source            {wb_path}")
    print("field counts      " + json.dumps({k: len(v) for k, v in tables["records"].items()}))
    print("annexure rows     " + json.dumps(
        {k: len(extract_rows(wb[SHEET], *v)) - 1 for k, v in ANNEXURES.items()}))
    print("state codes       " + str(len(assets["state_codes"])))
    print("section codes     " + str(len(assets["sections"])))
    print("reason flags      " + str(len(assets["reason_codes"])))
    for w in warnings:
        print(f"WARNING  {w}")
    if not warnings:
        print("shape check       all record and annexure counts match the expected edition")
    print(f"wrote {ASSETS / 'annexures.json'}")
    print(f"wrote {REFERENCES / 'form-140-layout.md'}")
    print(f"wrote {REFERENCES / 'section-codes.md'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
