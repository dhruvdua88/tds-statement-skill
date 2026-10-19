"""The working paper. This is the file a reviewer signs off, not the TXT.

A statement TXT is unreadable by design, so the deliverable that actually gets
reviewed is this workbook. It follows the house layout: a plain-language
summary, the challan allocation with the money that is still unallocated, the
deductee-to-challan bridge, the full register, the section map, everything that
was adjusted or flagged, what is still pending, and exactly what the utility
said.

The bridge sheet is the one that earns its place. It lists every deductee row
under the challan it settles, with a running sum and the balance left on that
challan, so an allocation error is visible as a red remaining figure rather
than as a rejection at upload.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path

import codes
import queries
from model import Statement

HEAD_FILL = "1F3864"
BAND_FILL = "F2F2F2"
OK_FILL = "E2EFDA"
BAD_FILL = "FCE4E4"
WARN_FILL = "FFF2CC"
MONEY = "#,##0.00"


def _style_header(ws, row: int = 1) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _widths(ws, widths: list[int]) -> None:
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build(st: Statement, out_path: str | Path, fvu_result: dict | None = None,
          issues: list[dict] | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    totals = st.totals
    # One query list feeds the Summary's blocker block, the Queries tab and the
    # standalone file package.py writes, so all three always tell the same story.
    qrows = queries.build(st, issues, fvu_result)
    qblockers = queries.blockers(qrows)

    # ---------------------------------------------------------------- summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "TDS statement working paper"
    ws["A1"].font = Font(bold=True, size=14)
    form_name = {"140": "Form 140 - resident non-salary (earlier Form 26Q)",
                 "138": "Form 138 - salary (earlier Form 24Q)",
                 "144": "Form 144 - payments to non-residents (earlier Form 27Q)",
                 "143": "Form 143 - TCS (earlier Form 27EQ)"}.get(st.form, st.form)
    ded = st.deductor
    addr = ded.get("address")
    rp = st.responsible_person
    rows = [
        ("Form", form_name),
        ("Statement type", "Original" if st.upload_type == "R" else "Correction"),
        ("Financial year / quarter", f"{st.financial_year} {st.quarter}"),
        ("Tax year / assessment year", f"{st.tax_year} / {st.assessment_year}"),
        ("TAN", st.tan),
        ("Deductor PAN", ded.get("pan", "")),
        ("Deductor name", ded.get("name", "")),
        ("Deductor category", f"{ded.get('category', '')} - {codes.DEDUCTOR_CATEGORIES.get(str(ded.get('category','')).upper(), '')}".strip(" -")),
        ("Address on the statement", ", ".join(x for x in [
            getattr(addr, "line1", ""), getattr(addr, "line2", ""), getattr(addr, "line3", ""),
            getattr(addr, "line4", ""), getattr(addr, "line5", ""), getattr(addr, "pin", "")] if x)),
        ("State code", f"{getattr(addr, 'state_code', '')} {codes.state_name(str(getattr(addr, 'state_code', '')))}".strip()),
        ("Deductor e-mail / phone", f"{ded.get('email', '')}  {ded.get('phone', '')}".strip()),
        ("Responsible person", f"{rp.get('name', '')}, {rp.get('designation', '')}".strip(", ")),
        ("Responsible person PAN / phone", f"{rp.get('pan', '') or '(blank)'}  {rp.get('phone', '') or '(blank)'}".strip()),
        ("GSTIN", ded.get("gstin", "")),
        ("", ""),
        ("Challans", totals["challans"]),
        ("Deductee rows", totals["rows"]),
        ("Tax (TDS)", totals["tax"]),
        ("Surcharge", totals["surcharge"]),
        ("Cess", totals["cess"]),
        ("Interest", totals["interest"]),
        ("Fee / penalty", totals["fee"] + totals["penalty"]),
        ("Total deposited", totals["deposit"]),
        ("Deductee tax in the statement", totals["deductee_tds"]),
        ("Unallocated / variance", totals["variance"]),
        ("", ""),
        ("Status", "DRAFT - do not upload" if st.is_draft else "Ready for the official FVU"),
        ("Outstanding", st.draft_reason or "nothing"),
        ("", ""),
        ("FVU", queries.fvu_status_label(fvu_result)),
        ("Blockers stopping the FVU",
         f"{len(qblockers)} - fix these before the utility will produce a .fvu"
         if qblockers else "none"),
        ("Prepared", datetime.now().strftime("%d-%b-%Y")),
    ]
    for q in qblockers:
        rows.append((f"   {q['id']}  {q['reference']}", q["issue"]))
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=v)
        if isinstance(v, Decimal):
            cell.number_format = MONEY
        if k in ("Unallocated / variance",) and isinstance(v, Decimal) and v != 0:
            cell.fill = PatternFill("solid", fgColor=BAD_FILL)
    ws["B5"].number_format = "@"
    _widths(ws, [34, 78])

    # ------------------------------------------------------- challan allocation
    ws = wb.create_sheet("Challan Allocation")
    ws.append(["No", "Challan serial", "Date", "BSR", "Minor head", "Tax",
               "Interest", "Total deposited", "Rows", "Allocated", "Remaining", "Status"])
    _style_header(ws)
    for row in st.allocation():
        ch = next((c for c in st.challans if c.no == row["no"]), None)
        ws.append([row["no"], row["serial"], row["date"], ch.bsr_code if ch else "",
                   ch.minor_head if ch else "", float(row["challan_tax"]), float(ch.interest) if ch else 0.0,
                   float(row["deposited"]), row["rows"], float(row["allocated"]),
                   float(row["remaining"]), "Ties" if row["remaining"] == 0 else "Open"])
        last = ws[ws.max_row]
        for c in last[5:6] + last[7:11]:
            c.number_format = MONEY
        last[11].fill = PatternFill("solid", fgColor=OK_FILL if row["remaining"] == 0 else BAD_FILL)
    tot = ws.max_row + 1
    ws.cell(row=tot, column=6, value=float(totals["tax"])).number_format = MONEY
    ws.cell(row=tot, column=8, value=float(totals["deposit"])).number_format = MONEY
    ws.cell(row=tot, column=10, value=float(sum(r["allocated"] for r in st.allocation()))).number_format = MONEY
    for c in ws[tot]:
        c.font = Font(bold=True)
    _widths(ws, [5, 15, 12, 12, 11, 14, 12, 16, 7, 14, 13, 10])
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------ bridge
    ws = wb.create_sheet("Deductee-Challan Bridge")
    ws.append(["Challan", "Challan date", "Row", "Deductee", "PAN", "Section",
               "TDS", "Running total", "Challan tax", "Remaining", "Status"])
    _style_header(ws)
    for row in st.allocation():
        ch_rows = [r for r in st.deductees if r.challan_no == row["no"]]
        running = Decimal("0.00")
        for i, r in enumerate(ch_rows, start=1):
            running += r.tds
            remaining = row["challan_tax"] - running
            ws.append([f"{row['no']} ({row['serial']})", row["date"], i, r.name, r.pan,
                       f"{r.section} {codes.section_label(r.section)}", float(r.tds),
                       float(running), float(row["challan_tax"]), float(remaining),
                       "Over" if remaining < 0 else "OK"])
            last = ws[ws.max_row]
            for c in last[6:10]:
                c.number_format = MONEY
            last[10].fill = PatternFill("solid", fgColor=BAD_FILL if remaining < 0 else OK_FILL)
    _widths(ws, [20, 12, 6, 40, 14, 34, 14, 14, 13, 13, 9])
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------ register
    ws = wb.create_sheet("Deductee Register")
    ws.append(["#", "Challan", "Deductee", "PAN", "Code", "Section", "Nature",
               "Payment date", "Deduction date", "Amount paid", "TDS", "Rate", "Status",
               "Remarks", "Source"])
    _style_header(ws)
    for i, r in enumerate(st.deductees, start=1):
        status = r.reason_code or ("deposited" if r.deposited == "Y" else "not deposited")
        ws.append([i, r.challan_no, r.name, r.pan, r.deductee_code, r.section,
                   codes.section_label(r.section), r.payment_date, r.deduction_date,
                   float(r.amount_paid), float(r.tds), r.rate, status, r.remarks, r.source_ref])
        last = ws[ws.max_row]
        last[9].number_format = MONEY
        last[10].number_format = MONEY
        if i % 2 == 0:
            for c in last:
                c.fill = PatternFill("solid", fgColor=BAND_FILL)
    _widths(ws, [5, 9, 40, 14, 6, 9, 34, 12, 13, 16, 14, 10, 7, 26, 16])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ------------------------------------------------------------ section map
    ws = wb.create_sheet("Section Map")
    ws.append(["Section", "Nature of payment", "Rate assumption", "Rows", "Amount paid", "TDS"])
    _style_header(ws)
    for code in sorted({r.section for r in st.deductees}):
        rows_ = [r for r in st.deductees if r.section == code]
        meta = codes.SECTIONS.get(code, {})
        rate = codes.indicative_rate(code)
        ws.append([code, meta.get("description", "(not in the table)"),
                   f"{rate}%" if rate is not None else "", len(rows_),
                   float(sum((r.amount_paid for r in rows_), Decimal("0.00"))),
                   float(sum((r.tds for r in rows_), Decimal("0.00")))])
        ws[ws.max_row][4].number_format = MONEY
        ws[ws.max_row][5].number_format = MONEY
    _widths(ws, [10, 40, 15, 7, 18, 16])

    # ------------------------------------------------------------ flags
    ws = wb.create_sheet("Adjustments & Flags")
    ws.append(["Level", "Code", "Detail", "Rows / reference"])
    _style_header(ws)
    flag_rows = [(f.get("level", "info"), f.get("code", ""), f.get("message", ""),
                  ", ".join(str(x) for x in f.get("rows", []) or [])) for f in st.flags]
    for i in issues or []:
        if i.get("level") in ("warn", "error"):
            flag_rows.append((i["level"], i.get("code", ""), i.get("message", ""), i.get("ref", "")))
    for level, code, msg, ref in flag_rows:
        ws.append([level, code, msg, ref])
        fill = {"error": BAD_FILL, "warn": WARN_FILL}.get(level)
        if fill:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=fill)
    _widths(ws, [8, 22, 96, 30])

    # ------------------------------------------------------------ queries
    ws = wb.create_sheet("Queries & Blockers")
    ws.append(["#", "Severity", "Source", "Where", "Issue / ask", "Why it matters", "Owner", "Status"])
    _style_header(ws)
    sev_fill = {"BLOCKER": BAD_FILL, "WARNING": WARN_FILL, "INFO": OK_FILL}
    for q in qrows:
        ws.append([q["id"], q["severity"], queries.SOURCE_LABEL.get(q["source"], q["source"]),
                   q["reference"], q["issue"], q["why"], q["owner"], q["status"]])
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=sev_fill.get(q["severity"], WARN_FILL))
    if not qrows:
        ws.append(["", "", "", "", "Nothing outstanding.", "", "", ""])
    _widths(ws, [6, 10, 22, 34, 58, 58, 9, 9])
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------ FVU
    ws = wb.create_sheet("FVU Results")
    ws.append(["Line", "Record", "Field", "Challan", "Deductee", "Code", "Message"])
    _style_header(ws)
    if fvu_result:
        for e in fvu_result.get("errors", []):
            ws.append([e.get("line_no", ""), e.get("record_type", ""), e.get("field_name", ""),
                       e.get("challan_no", ""), e.get("deductee_no", ""),
                       e.get("code", ""), e.get("message", e.get("raw", ""))])
        ws.append([])
        for k in ("status", "mode", "exit_code", "jar", "format_version", "csi", "fvu", "input"):
            ws.append([k, fvu_result.get(k, "")])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    else:
        ws.append(["", "", "", "", "", "", "The official utility has not been run yet."])
    _widths(ws, [8, 10, 34, 12, 10, 14, 90])

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
