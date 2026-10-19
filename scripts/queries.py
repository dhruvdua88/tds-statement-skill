"""One list of everything standing between this statement and a filed return.

A statement is not finished when the TXT is written; it is finished when the
utility stops complaining. Until now the three things that can complain each
spoke only to whoever happened to read its own output: the decisions the builder
could not make, the pre-flight checks, and the utility's own error file. A
reviewer had to know to look in three places, and the client got whichever of
them was pasted into an e-mail.

This module merges them into one list. Every entry names the field or record it
is about, states the issue, says why it matters, and says who has to act. That
one list is the workbook's Queries tab, the standalone query file, and the
blockers section of the README and the upload checklist.

Severity is narrow on purpose:

* ``BLOCKER``  no ``.fvu`` can exist until this is resolved, so nothing may be
  uploaded. These are the "blockers for generating the FVU".
* ``WARNING``  the statement is filable but something needs a human eye.
* ``INFO``     recorded so the workpaper is complete.

Nothing here decides whether the statement is correct. It reports what is
outstanding, and the only all-clear it will give is a ``.fvu`` in hand.
"""

from __future__ import annotations

from pathlib import Path

BLOCKER = "BLOCKER"
WARNING = "WARNING"
INFO = "INFO"

# Blocker first, so the list reads top-down as "fix these before anything else".
SEVERITY_RANK = {BLOCKER: 0, WARNING: 1, INFO: 2}
SEVERITY_LABEL = {
    BLOCKER: "Blocker - stops the .fvu",
    WARNING: "Warning - needs a human eye",
    INFO: "Recorded",
}
SOURCE_LABEL = {
    "intake": "Builder could not decide",
    "lint": "Pre-flight check",
    "fvu": "Official utility",
    "csi": "Challan Status Inquiry",
}

# The linter's own levels are less specific than ours: everything it calls an
# error is a blocker, because the utility is expected to reject the file.
LINT_SEVERITY = {"error": BLOCKER, "warn": WARNING, "info": INFO}


def _row(severity: str, source: str, reference: str, issue: str, why: str,
         owner: str = "", status: str = "open") -> dict:
    return {
        "severity": severity,
        "source": source,
        "reference": reference.strip() or "(the statement)",
        "issue": issue.strip(),
        "why": why.strip(),
        "owner": owner.strip().lower(),
        "status": status,
    }


# ---------------------------------------------------------------------------
# The three sources
# ---------------------------------------------------------------------------
def from_pending(st) -> list[dict]:
    """Decisions the builder needed and did not have (the document's ``pending``).

    An entry may carry its own ``question`` (what to put to the client) and
    ``severity``. Otherwise the field label becomes the issue and the entry is a
    blocker, which is the safe default: an unanswered question is not a filing.
    """
    out = []
    for p in st.pending:
        blocking = bool(p.get("blocking", True))
        field = str(p.get("field", "") or "").strip() or "(unspecified)"
        severity = str(p.get("severity", "") or "").upper()
        if severity not in SEVERITY_RANK:
            severity = BLOCKER if blocking else WARNING
        issue = str(p.get("question", "") or "").strip() or f"{field} is needed but not on any record we hold."
        out.append(_row(severity, "intake", field, issue,
                        str(p.get("why", "") or ""), str(p.get("owner", "") or "")))
    return out


# The linter restates the document's own pending items as PENDING errors so a
# bare ``tds lint`` shows them without reading the document. They are the same
# items ``from_pending`` already carries, with a better reference, so the
# restatements are dropped here to keep one row per problem.
RESTATED_BY_INTAKE = {"PENDING"}


def from_lint(issues: list[dict] | None) -> list[dict]:
    out = []
    for i in issues or []:
        if str(i.get("code", "")).upper() in RESTATED_BY_INTAKE:
            continue
        level = str(i.get("level", "info") or "info").lower()
        out.append(_row(
            LINT_SEVERITY.get(level, INFO), "lint",
            str(i.get("ref", "") or "(the statement)"),
            f"{i.get('code', '')}: {i.get('message', '')}".strip(": "),
            "A pre-flight finding. The utility is expected to reject or query this, "
            "so it has to be cleared before the run.",
            "us",
        ))
    return out


def from_fvu(st, fvu_result: dict | None) -> list[dict]:
    """What the official utility said, and the one thing that stops it short.

    A real Challan Status Inquiry file is a hard requirement, not an
    inconvenience: the utility matches every challan against it, and without it
    there is no ``.fvu``. It is a blocker in its own right, so it is listed even
    when the field-level checks are clean.
    """
    out = []
    status = (fvu_result or {}).get("status", "NOT_RUN")

    if status == "PASS":
        out.append(_row(INFO, "fvu", "the official utility",
                        "The utility produced a .fvu for this statement.",
                        "This is the only proof the statement validates, and it is the "
                        "file the portal accepts.", "us", "done"))
        return out

    if fvu_result is None:
        out.append(_row(BLOCKER, "fvu", "the official utility has not been run",
                        "Run the File Validation Utility over the built TXT.",
                        "A pre-flight pass is not proof. Only the utility writes a .fvu, "
                        "and only a .fvu may be uploaded.", "us"))
    else:
        for e in fvu_result.get("errors", []) or []:
            ref = " ".join(x for x in [
                f"line {e.get('line_no', '')}".strip() if e.get("line_no") else "",
                str(e.get("record_type", "") or ""),
                f"f{e.get('field_name')}" if e.get("field_name") else "",
            ] if x) or "(whole statement)"
            out.append(_row(
                BLOCKER, "fvu", ref,
                f"{e.get('code', '?')}: {e.get('message', e.get('raw', ''))}".strip(),
                e.get("meaning", "") or "The utility rejected this record.", "us"))
        if status == "STOPPED":
            out.append(_row(BLOCKER, "fvu", "the utility stopped",
                            "The utility ran but wrote neither a .fvu nor an error file.",
                            "Usually a placeholder CSI, or a run that did not finish. "
                            "Check the log beside the jar.", "us"))

    if not (fvu_result or {}).get("csi"):
        out.append(_row(
            BLOCKER, "csi",
            f"Challan Status Inquiry (.csi) for TAN {st.tan or '(not set)'}",
            "Get the real CSI for this TAN and quarter from the deductor's e-filing "
            "login (Challan Status Inquiry) and read the six challan digests out of it.",
            "The utility matches every challan against the CSI. Without the real file it "
            "cannot produce a .fvu - it stops at T-FV-3175 or fails outright.", "client"))
    return out


# ---------------------------------------------------------------------------
# The list
# ---------------------------------------------------------------------------
def build(st, issues: list[dict] | None = None, fvu_result: dict | None = None) -> list[dict]:
    rows = from_pending(st) + from_lint(issues) + from_fvu(st, fvu_result)
    # Stable sort: blockers first, everything else in the order it was found.
    rows.sort(key=lambda q: SEVERITY_RANK.get(q["severity"], 9))
    for n, q in enumerate(rows, 1):
        q["id"] = f"Q-{n:02d}"
    return rows


def blockers(rows: list[dict]) -> list[dict]:
    return [q for q in rows if q["severity"] == BLOCKER]


def counts(rows: list[dict]) -> dict:
    return {sev: sum(1 for q in rows if q["severity"] == sev) for sev in SEVERITY_RANK}


def ready(rows: list[dict]) -> bool:
    """A statement is ready only when nothing blocks the utility."""
    return not blockers(rows)


def fvu_status_label(fvu_result: dict | None) -> str:
    if not fvu_result:
        return "BLOCKED - the utility has not been run"
    status = fvu_result.get("status", "NOT_RUN")
    if status == "PASS":
        return "PASS - a .fvu was produced"
    if status == "FIELD_ERRORS":
        return f"BLOCKED - {fvu_result.get('error_count', len(fvu_result.get('errors', [])))} field error(s)"
    if status == "STOPPED_NO_CSI":
        return "BLOCKED - no real CSI was supplied"
    return f"BLOCKED - the run did not finish ({status})"


# ---------------------------------------------------------------------------
# The standalone query list
# ---------------------------------------------------------------------------
def to_markdown(st, rows: list[dict], fvu_result: dict | None = None) -> str:
    n = counts(rows)
    who = str(st.deductor.get("name", "") or "").strip() or "(deductor name not supplied yet)"
    title = (f"# Queries and blockers - {who} - Form {st.form}, "
             f"{st.financial_year} {st.quarter}")
    out = [
        title, "",
        f"TAN {st.tan or '(not set)'}",
        f"FVU: {fvu_status_label(fvu_result)}",
        f"Blockers {n[BLOCKER]}  |  Warnings {n[WARNING]}  |  Recorded {n[INFO]}",
        "",
    ]
    bl = blockers(rows)
    if not bl:
        out += ["Nothing blocks the utility. The `.fvu` is the proof; this list is empty of blockers.", ""]

    for sev in (BLOCKER, WARNING, INFO):
        group = [q for q in rows if q["severity"] == sev]
        if not group:
            continue
        out += [f"## {SEVERITY_LABEL[sev]} ({len(group)})", "",
                "| # | Where | Issue | Why it matters | Owner |",
                "|---|---|---|---|---|"]
        for q in group:
            out.append("| {} | {} | {} | {} | {} |".format(
                q["id"], q["reference"], q["issue"].replace("|", "\\|"),
                q["why"].replace("|", "\\|"), q["owner"] or "-"))
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def write_markdown(st, path: str | Path, issues: list[dict] | None = None,
                   fvu_result: dict | None = None) -> Path:
    rows = build(st, issues, fvu_result)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(to_markdown(st, rows, fvu_result), encoding="utf-8")
    return out


def write_xlsx(st, path: str | Path, issues: list[dict] | None = None,
               fvu_result: dict | None = None) -> Path:
    """A one-sheet workbook: the query list on its own, for sending to the client."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    head_fill = PatternFill("solid", fgColor="1F3864")
    fills = {BLOCKER: PatternFill("solid", fgColor="FCE4E4"),
             WARNING: PatternFill("solid", fgColor="FFF2CC"),
             INFO: PatternFill("solid", fgColor="E2EFDA")}
    rows = build(st, issues, fvu_result)
    n = counts(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Queries"
    ws["A1"] = f"Queries and blockers - Form {st.form}, {st.financial_year} {st.quarter}, TAN {st.tan or '(not set)'}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"FVU: {fvu_status_label(fvu_result)}"
    ws["A3"] = f"Blockers {n[BLOCKER]}  |  Warnings {n[WARNING]}  |  Recorded {n[INFO]}"
    ws["A4"] = "Blockers stop the .fvu. Nothing is uploaded until this list is empty of them."

    header = ["#", "Severity", "Where", "Issue / ask", "Why it matters", "Owner", "Status"]
    ws.append([])
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    for q in rows:
        ws.append([q["id"], q["severity"], q["reference"], q["issue"], q["why"],
                   q["owner"], q["status"]])
        last = ws[ws.max_row]
        for c in last:
            c.fill = fills[q["severity"]]
        last[3].alignment = Alignment(wrap_text=True, vertical="top")
        last[4].alignment = Alignment(wrap_text=True, vertical="top")
    if not rows:
        ws.append(["", "", "", "Nothing outstanding.", "", "", ""])
    for col, width in zip("ABCDEFG", [6, 10, 34, 60, 60, 9, 9]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A6"

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def write_sidecar(st, out_dir: str | Path, issues: list[dict] | None = None,
                  fvu_result: dict | None = None) -> list[Path]:
    """Write both standalone forms of the list into a folder."""
    out = Path(out_dir)
    name = f"Queries_{st.form}_{(st.quarter or 'Q0')}"
    return [
        write_xlsx(st, out / f"{name}.xlsx", issues, fvu_result),
        write_markdown(st, out / f"{name}.md", issues, fvu_result),
    ]
