"""Books to return: section-wise TDS per the ledger, per the statement and per the challans.

The input is a ledger export from Tally or Zoho, as xlsx or csv. Two kinds:

* ``tds``      the TDS payable ledger(s). Credits are tax deducted; debits are
               deposits and reversals. The books figure is the credits.
* ``expense``  an expense ledger. Debits are the amounts booked; the books
               figure is compared with the amount paid or credited on the
               statement rather than with the tax.

The section of a ledger row comes from a section column if there is one,
otherwise from the ledger name (a "Ledger" column, or the title rows above the
table), read the way intake reads a section: "TDS Payable - 194C" and
"Contractor 393(1)" both resolve. A row whose section cannot be read is listed,
never dropped.

**Sign conventions.** Tally often exports with the signs inverted: a credit
comes out negative in a single Amount column. So the side of an entry is read,
in order of preference, from separate Debit and Credit columns, from a Dr/Cr
marker (its own column or a suffix on the amount), and only then from the
sign of a single amount. For that last case the caller says which sign means
credit (``credit_sign``, default "negative", the Tally habit) because nothing
in a bare signed number can tell the two apart. Amounts on each side are taken
as absolute values, so an export that writes credits as positive in a Credit
column and one that writes them as negative both come out the same.
"""

from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

import intake
from model import ModelError, Statement

ZERO = Decimal("0.00")
HEADERS = {
    "debit": ["debit", "dr", "debit amount", "withdrawal"],
    "credit": ["credit", "cr", "credit amount", "deposit"],
    "amount": ["amount", "value", "net amount"],
    "drcr": ["dr/cr", "dr cr", "type", "drcr"],
    "ledger": ["ledger", "ledger name", "account", "account name"],
    "section": ["section", "tds section", "section code", "nature of payment"],
    "date": ["date", "voucher date", "transaction date"],
    "particulars": ["particulars", "party", "description", "narration", "name"],
    "voucher": ["voucher no", "vch no", "voucher number", "reference", "ref no"],
}
UNMAPPED = "(section not read)"


class ReconError(ModelError):
    pass


def _norm(text) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9/ ]", " ", str(text or "").lower())).strip()


def _columns(row: list) -> dict[str, int]:
    found: dict[str, int] = {}
    for i, cell in enumerate(row):
        h = _norm(cell)
        for key, names in HEADERS.items():
            if key not in found and h in names:
                found[key] = i
    return found


def _side(raw, drcr: str, credit_sign: str) -> tuple[str, Decimal]:
    """('dr'|'cr', absolute amount) for a single-amount row."""
    text = str(raw or "").strip()
    marker = (drcr or "").strip().lower()
    m = re.search(r"\b(dr|cr)\.?\s*$", text, re.I)
    if m and not marker:
        marker = m.group(1).lower()
    value = intake.cell_money(text if not isinstance(raw, (int, float, Decimal)) else raw)
    if marker.startswith("cr"):
        return "cr", abs(value)
    if marker.startswith("dr"):
        return "dr", abs(value)
    negative_is_credit = credit_sign == "negative"
    if value < 0:
        return ("cr" if negative_is_credit else "dr"), abs(value)
    return ("dr" if negative_is_credit else "cr"), value


def read_ledger(path: str | Path, credit_sign: str = "negative") -> list[dict]:
    if credit_sign not in ("negative", "positive"):
        raise ReconError("credit_sign must be 'negative' or 'positive'")
    entries = []
    for sheet, rows in intake.read_tables(path).items():
        head, cols = -1, {}
        for i, row in enumerate(rows[:40]):
            c = _columns(row)
            if ("debit" in c and "credit" in c) or "amount" in c:
                head, cols = i, c
                break
        if head < 0:
            continue
        title = " ".join(str(c) for r in rows[:head] for c in r if c not in (None, ""))
        title_section = intake.cell_section(title)
        for n, row in enumerate(rows[head + 1:], start=head + 2):
            get = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None  # noqa: E731
            first = _norm(row[0] if row else "")
            label = _norm(get("particulars"))
            if not any(str(c or "").strip() for c in row) or first in intake.TOTAL_WORDS \
                    or label in intake.TOTAL_WORDS \
                    or any(x.startswith(("opening", "closing")) for x in (first, label)):
                continue
            if "debit" in cols and "credit" in cols:
                dr, cr = abs(intake.cell_money(get("debit"))), abs(intake.cell_money(get("credit")))
            else:
                side, amt = _side(get("amount"), str(get("drcr") or ""), credit_sign)
                dr, cr = (amt, ZERO) if side == "dr" else (ZERO, amt)
            if not dr and not cr:
                continue
            section = (intake.cell_section(get("section")) or intake.cell_section(get("ledger"))
                       or title_section or UNMAPPED)
            entries.append({
                "ref": f"{Path(path).name}:{sheet}!{n}", "date": intake.cell_date(get("date")) if get("date") else "",
                "particulars": str(get("particulars") or get("ledger") or "").strip(),
                "voucher": str(get("voucher") or "").strip(), "section": section,
                "debit": dr, "credit": cr,
            })
    if not entries:
        raise ReconError(f"{Path(path).name}: no ledger table found (need Debit and Credit "
                         "columns, or an Amount column)")
    return entries


def challan_by_section(st: Statement) -> dict[str, Decimal]:
    """Challan tax by section: the challan's own section if recorded, else its rows'.

    A challan's tax not taken up by any row is reported against no section
    rather than spread by guesswork.
    """
    out: dict[str, Decimal] = {}
    for ch in st.challans:
        if ch.tds_section:
            out[ch.tds_section] = out.get(ch.tds_section, ZERO) + ch.all_tax
            continue
        used = ZERO
        for r in st.deductees:
            if r.challan_no == ch.no:
                out[r.section] = out.get(r.section, ZERO) + r.total_tds
                used += r.total_tds
        if ch.all_tax != used:
            key = f"(unallocated, challan {ch.challan_serial or ch.no})"
            out[key] = out.get(key, ZERO) + ch.all_tax - used
    return out


def reconcile(st: Statement, entries: list[dict], kind: str = "tds") -> dict:
    if kind not in ("tds", "expense"):
        raise ReconError("ledger kind must be 'tds' or 'expense'")
    books: dict[str, Decimal] = {}
    for e in entries:
        amount = e["credit"] if kind == "tds" else e["debit"] - e["credit"]
        books[e["section"]] = books.get(e["section"], ZERO) + amount
    statement: dict[str, Decimal] = {}
    for r in st.deductees:
        statement[r.section] = statement.get(r.section, ZERO) + (
            r.total_tds if kind == "tds" else r.amount_paid)
    challans = challan_by_section(st) if kind == "tds" else {}
    rows = []
    for sec in sorted(set(books) | set(statement) | set(challans)):
        b, s, c = books.get(sec, ZERO), statement.get(sec, ZERO), challans.get(sec, ZERO)
        rows.append({"section": sec, "books": b, "statement": s, "challans": c if kind == "tds" else None,
                     "books_less_statement": b - s,
                     "statement_less_challans": (s - c) if kind == "tds" else None})
    differences = [r for r in rows if r["books_less_statement"] or (r["statement_less_challans"] or ZERO)]
    return {"kind": kind, "rows": rows, "differences": differences,
            "unmapped": [e for e in entries if e["section"] == UNMAPPED],
            "debits": sum((e["debit"] for e in entries), ZERO),
            "credits": sum((e["credit"] for e in entries), ZERO)}


def write_sheet(wb, st: Statement, result: dict) -> None:
    from openpyxl.styles import Font, PatternFill

    import codes
    import excel

    tds = result["kind"] == "tds"
    ws = wb.create_sheet("Books Reconciliation")
    ws["A1"] = ("TDS per books, per statement and per challans" if tds
                else "Amount booked to expense vs amount paid/credited on the statement")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"Ledger debits {result['debits']}, credits {result['credits']}. "
                + ("Books = credits to the TDS ledger (tax deducted)." if tds
                   else "Books = debits less credits to the expense ledger."))
    head = (["Section", "Nature", "Per books", "Per statement", "Per challans",
             "Books - statement", "Statement - challans"] if tds else
            ["Section", "Nature", "Per books", "Per statement", "Books - statement"])
    ws.append([])
    ws.append(head)
    excel._style_header(ws, ws.max_row)
    for r in result["rows"]:
        vals = [r["section"], codes.section_label(r["section"]) if r["section"][:1].isdigit() else "",
                float(r["books"]), float(r["statement"])]
        vals += [float(r["challans"]), float(r["books_less_statement"]), float(r["statement_less_challans"])] \
            if tds else [float(r["books_less_statement"])]
        ws.append(vals)
        last = ws[ws.max_row]
        for c in last[2:]:
            c.number_format = excel.MONEY
        if r in result["differences"]:
            for c in last:
                c.fill = PatternFill("solid", fgColor=excel.WARN_FILL)

    ws.append([])
    ws.append(["Differences, one row each"])
    ws[ws.max_row][0].font = Font(bold=True)
    if not result["differences"] and not result["unmapped"]:
        ws.append(["None. Books, statement" + (" and challans" if tds else "") + " agree section by section."])
    for r in result["differences"]:
        if r["books_less_statement"]:
            ws.append([r["section"], f"books {r['books']} vs statement {r['statement']}: "
                       f"difference {r['books_less_statement']}"])
        if tds and r["statement_less_challans"]:
            ws.append([r["section"], f"statement {r['statement']} vs challans {r['challans']}: "
                       f"difference {r['statement_less_challans']}"])
    for e in result["unmapped"]:
        ws.append([UNMAPPED, f"{e['ref']} {e['date']} {e['particulars']} {e['voucher']} "
                             f"Dr {e['debit']} Cr {e['credit']}".strip()])
    excel._widths(ws, [30, 70, 16, 16, 16, 18, 18])


def write_workbook(st: Statement, result: dict, out: str | Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, st, result)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
